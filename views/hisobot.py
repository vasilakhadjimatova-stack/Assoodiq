from datetime import date
from io import BytesIO

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core import finance as F
from core.models import Program, Transaction
from core.extensions import db

bp = Blueprint("hisobot", __name__, url_prefix="/hisobot")


def _year():
    try:
        return int(request.args.get("yil", date.today().year))
    except ValueError:
        return date.today().year


def _years():
    first = db.session.query(db.func.min(Transaction.date)).scalar() or date.today()
    return list(range(date.today().year, first.year - 1, -1))


@bp.route("/")
def pnl():
    year = _year()
    paid_map = F.enrollment_paid_map()
    programs = [F.program_stats(p, paid_map) for p in Program.query.all()]
    programs.sort(key=lambda s: -s["profit"])
    return render_template("pages/pnl.html", year=year, years=_years(), t=F.pnl_table(year),
                           months=F.MONTHS_SHORT, programs=programs)


@bp.route("/excel")
def pnl_excel():
    year = _year()
    t = F.pnl_table(year)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Foyda-zarar {year}"
    head = ["Modda"] + F.MONTHS_FULL + ["Jami"]
    ws.append([f"{F.get_setting('org_name')} — foyda va zarar hisoboti, {year} (so'm)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(head)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E8F0EC")
    for c in ws[3]:
        c.font, c.fill = bold, fill

    def section(title, sec):
        ws.append([title])
        ws.cell(ws.max_row, 1).font = bold
        for g in sec["groups"]:
            ws.append(["  " + g["name"]] + [round(v) for v in g["months"]] + [round(g["total"])])
        ws.append([f"Jami: {title.lower()}"] + [round(v) for v in sec["totals"]] + [round(sec["total"])])
        for c in ws[ws.max_row]:
            c.font = bold

    section("Daromadlar", t["income"])
    ws.append([])
    section("Xarajatlar", t["expense"])
    ws.append([])
    ws.append(["Sof foyda"] + [round(v) for v in t["profit"]] + [round(t["profit_total"])])
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, fill

    ws.column_dimensions["A"].width = 32
    for col in range(2, 15):
        ws.column_dimensions[ws.cell(3, col).column_letter].width = 14
    for row in ws.iter_rows(min_row=4, min_col=2):
        for c in row:
            c.number_format = "# ##0"
            c.alignment = Alignment(horizontal="right")
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"assoodiq-foyda-zarar-{year}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

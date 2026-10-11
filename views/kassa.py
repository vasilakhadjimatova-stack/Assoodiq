from datetime import date
from io import BytesIO

from flask import Blueprint, flash, g, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, or_

from core import audit, sheets
from core import finance as F
from core.extensions import db
from core.models import (
    AQD_TYPES, SHARIAT, Account, Category, Enrollment, Program, Transaction,
)
from views import form_date, form_int, money, safe_next

bp = Blueprint("kassa", __name__, url_prefix="/kassa")

KIND_FILTERS = {
    "": "Hammasi",
    "kirim": "Kirim",
    "chiqim": "Chiqim",
    "otkazma": "O'tkazma",
    "qarz": "Qarz",
}


def _month_arg():
    raw = request.args.get("oy", "")
    try:
        y, m = (int(x) for x in raw.split("-"))
        return date(y, m, 1)
    except ValueError:
        return date.today().replace(day=1)


def _filters():
    """URL'dagi filtrlar: oy, tur, hisob, kategoriya, qidiruv (sahifa va Excel eksport uchun bir xil).

    period — faqat oy va hisob bo'yicha (qoldiq shundan hisoblanadi), base — barcha filtrlar bilan.
    """
    month = _month_arg()
    start, end = F.month_bounds(month)
    kind = request.args.get("tur", "")
    acc_id = form_int(request.args.get("hisob"))
    kat_raw = request.args.get("kat", "")
    no_cat = kat_raw == "yoq"  # kategoriyasiz kirim/chiqimlar (importdan keyin tekshirish uchun)
    cat_id = "yoq" if no_cat else form_int(kat_raw)
    q_text = (request.args.get("q") or "").strip()
    period = Transaction.query.filter(Transaction.date >= start, Transaction.date <= end)
    if acc_id:
        period = period.filter(or_(Transaction.account_id == acc_id, Transaction.to_account_id == acc_id))
    base = period
    if no_cat:
        base = base.filter(Transaction.category_id.is_(None), Transaction.kind.in_(("kirim", "chiqim")))
    elif cat_id:
        base = base.filter(Transaction.category_id == cat_id)
    if q_text:
        like = f"%{q_text}%"
        base = base.filter(or_(Transaction.counterparty.ilike(like), Transaction.note.ilike(like)))
    return {"month": month, "start": start, "end": end, "kind": kind, "acc_id": acc_id,
            "cat_id": cat_id, "q": q_text, "period": period, "base": base}


@bp.route("/")
def ledger():
    fl = _filters()
    month, start, end, kind, acc_id = fl["month"], fl["start"], fl["end"], fl["kind"], fl["acc_id"]
    base = fl["base"]

    counts = {}
    for k in KIND_FILTERS:
        counts[k] = _kind_filter(base, k).count()
    rows = _kind_filter(base, kind).order_by(Transaction.date.desc(), Transaction.id.desc()).all()

    # Daftar (jadval) ko'rinishi: eskidan yangiga, har qatordan keyingi qoldiq bilan
    account = db.session.get(Account, acc_id) if acc_id else None
    all_rows = fl["period"].order_by(Transaction.date, Transaction.id).all()
    opening, ledger_cur, flows = F.ledger_rows(all_rows, account=account, start=start)
    table = [dict(t=t, **flows[t.id]) for t in reversed(rows)]
    closing = flows[all_rows[-1].id]["bal"] if all_rows else opening
    sums = {
        "inc": sum(r["inc"] for r in table if not r["move"]),
        "out": sum(r["out"] for r in table if not r["move"]),
    }
    # Ro'yxat ko'rinishi: kunlik yakun (sana sarlavhasi yonida)
    day_totals = {}
    for r in table:
        if r["move"]:
            continue
        d = day_totals.setdefault(r["t"].date, [0.0, 0.0])
        d[0] += r["inc"]
        d[1] += r["out"]

    filtered = bool(kind or fl["cat_id"] or fl["q"])
    no_cat_count = fl["period"].filter(Transaction.category_id.is_(None),
                                       Transaction.kind.in_(("kirim", "chiqim"))).count()
    today = date.today()
    qe_date = today if start <= today <= end else (end if end < today else start)
    totals = F.period_totals(start, end)
    return render_template(
        "pages/kassa.html",
        rows=rows, kind=kind, kinds=KIND_FILTERS, counts=counts, acc_id=acc_id, q=fl["q"],
        cat_id=fl["cat_id"], account=account, filtered=filtered, no_cat_count=no_cat_count,
        table=table, opening=opening, qe_date=qe_date, closing=closing, sums=sums, ledger_cur=ledger_cur, day_totals=day_totals,
        month=month, prev_month=F.shift_month(month, -1), next_month=F.shift_month(month, 1),
        month_name=f"{F.MONTHS_FULL[month.month - 1]} {month.year}",
        totals=totals, **form_context(),
    )


def _kind_filter(q, kind):
    if kind == "qarz":
        return q.filter(Transaction.kind.in_(("qarz_kirim", "qarz_chiqim")))
    if kind:
        return q.filter(Transaction.kind == kind)
    return q


def form_context():
    """Kirim/chiqim oynasi uchun ro'yxatlar (boshqa sahifalar ham ishlatadi)."""
    hints = counterparty_hints()
    members = member_payment_hints()
    return {
        "f_accounts": F.accounts_with_balance(),
        "f_categories": Category.query.filter_by(active=True).order_by(Category.kind, Category.sort, Category.name).all(),
        "f_programs": Program.query.filter(Program.status != "yakunlangan").order_by(Program.name).all(),
        "f_hints": hints,
        "f_members": members,
        "f_names": sorted(set(hints) | {v["name"] for v in members.values()}, key=str.lower),
        "aqd_types": AQD_TYPES,
        "shariat_types": SHARIAT,
    }


def member_payment_hints():
    """{ism (kichik harfda): {name, items: [{e, p, pid, c, d}]}} — kassada o'quvchi ismi
    yozilganda kirimni uning dastur to'loviga bog'lash uchun (qarzidan ayiriladi)."""
    paid_map = F.enrollment_paid_map()
    pay_cat = {c.program_kind: c.id for c in Category.query.filter(Category.kind == "kirim",
                                                                     Category.program_kind.isnot(None),
                                                                     Category.active.is_(True)).all()}
    out = {}
    enrs = (Enrollment.query.filter(Enrollment.status != "chiqqan")
            .order_by(Enrollment.joined_at.desc()).all())
    for e in enrs:
        row = F.enrollment_row(e, paid_map)
        name = e.member.full_name.strip()
        rec = out.setdefault(name.lower(), {"name": name, "items": []})
        rec["items"].append({"e": e.id, "p": e.program.name, "pid": e.program_id,
                             "c": pay_cat.get(e.program.kind), "d": round(row["due"])})
    for rec in out.values():  # qarzi borlari birinchi
        rec["items"].sort(key=lambda i: i["d"] <= 0)
    return out


def counterparty_hints(limit=500):
    """Avval yozilgan "kimdan/kimga" nomlari va ularning oxirgi tur/kategoriya/hisob/dasturi.

    Ism tanlanganda qolgan maydonlar o'zi to'ladi — Sheets'dagi avto-to'ldirish kabi.
    """
    rows = (
        db.session.query(Transaction.counterparty, Transaction.kind, Transaction.category_id,
                         Transaction.account_id, Transaction.program_id)
        .filter(Transaction.counterparty != "", Transaction.counterparty.isnot(None),
                Transaction.kind.in_(("kirim", "chiqim")))
        .order_by(Transaction.date.desc(), Transaction.id.desc()).limit(limit).all()
    )
    hints = {}
    for name, kind, cat_id, acc_id, prog_id in rows:
        hints.setdefault(name.strip(), {"k": kind, "c": cat_id, "a": acc_id, "p": prog_id})
    return hints


def _parse(f, kind, account):
    """Formadan yozuv maydonlari. Xato bo'lsa (None, xabar) qaytaradi."""
    fields = {
        "date": form_date(f.get("date")),
        "note": (f.get("note") or "").strip()[:500],
        "counterparty": (f.get("counterparty") or "").strip()[:160],
    }
    if kind == "otkazma":
        to_acc = db.session.get(Account, form_int(f.get("to_account_id")) or 0)
        if to_acc is None or to_acc.id == account.id:
            return None, "O'tkazma uchun boshqa hisobni tanlang."
        to_amount = money(f.get("to_amount"))
        if not to_amount:
            to_amount = round(F.convert(money(f.get("amount")), account.currency, to_acc.currency), 2)
        fields.update(to_account_id=to_acc.id, to_amount=to_amount)
        return fields, None
    cat = db.session.get(Category, form_int(f.get("category_id")) or 0)
    fields["category_id"] = cat.id if cat and cat.kind == kind else None
    prog = db.session.get(Program, form_int(f.get("program_id")) or 0)
    fields["program_id"] = prog.id if prog else None
    enr = db.session.get(Enrollment, form_int(f.get("enrollment_id")) or 0)
    if enr and kind == "kirim":
        fields["enrollment_id"] = enr.id
        fields["program_id"] = enr.program_id
        if not fields["counterparty"]:
            fields["counterparty"] = enr.member.full_name
    if kind == "kirim":
        fields["aqd"] = f.get("aqd") if f.get("aqd") in AQD_TYPES else ""
        fields["shariat"] = f.get("shariat") if f.get("shariat") in SHARIAT else "halol"
    else:
        fields["aqd"], fields["shariat"] = "", "halol"
    return fields, None


@bp.route("/yangi", methods=["POST"])
def create():
    f = request.form
    if form_int(f.get("tx_id")):
        return update(form_int(f.get("tx_id")))
    kind = f.get("kind", "kirim")
    if g.user.role == "menejer" and kind != "kirim":
        flash("Menejer faqat kirim yozishi mumkin.", "bad")
        return _back()
    if kind not in ("kirim", "chiqim", "otkazma"):
        kind = "kirim"
    amount = money(f.get("amount"))
    account = db.session.get(Account, form_int(f.get("account_id")) or 0)
    if amount <= 0 or account is None:
        flash("Summa va hisobni ko'rsating.", "bad")
        return _back()
    fields, err = _parse(f, kind, account)
    if err:
        flash(err, "bad")
        return _back()

    tx = F.record(kind, amount, account, created_by=g.user.id, **fields)
    db.session.commit()
    _warn(tx)
    flash(f"{'Kirim' if kind == 'kirim' else 'Chiqim' if kind == 'chiqim' else 'O‘tkazma'} saqlandi: "
          f"{amount:,.0f} {'$' if account.currency == 'USD' else 'so‘m'}".replace(",", " "), "ok")
    sep = "&" if "?" in _back_url() else "?"
    if f.get("quick"):  # jadval ostidagi tez yozish qatori: fokus yana o'sha yerga qaytadi
        return redirect(_back_url() + sep + f"tez=1&belgi={tx.id}")
    if f.get("again"):
        return redirect(_back_url() + sep + f"yangi={kind}")
    return _back()


def update(tx_id):
    """Kirim/chiqimni tahrirlash. O'tkazma va qarz harakati o'chirib qayta yoziladi."""
    tx = db.get_or_404(Transaction, tx_id)
    if g.user.role == "menejer":
        flash("Tahrirlash uchun ruxsat yo'q.", "bad")
        return _back()
    f = request.form
    kind = f.get("kind") if f.get("kind") in ("kirim", "chiqim") else tx.kind
    if tx.kind not in ("kirim", "chiqim") or kind not in ("kirim", "chiqim"):
        flash("O'tkazma va qarz harakatini tahrirlab bo'lmaydi — o'chirib, qaytadan yozing.", "bad")
        return _back()
    amount = money(f.get("amount"))
    account = db.session.get(Account, form_int(f.get("account_id")) or 0)
    if amount <= 0 or account is None:
        flash("Summa va hisobni ko'rsating.", "bad")
        return _back()
    before = audit.describe_tx(tx)
    money_changed = (amount, account.id) != (tx.amount, tx.account_id)
    fields, _ = _parse(f, kind, account)
    if not fields.get("enrollment_id"):
        fields["enrollment_id"] = tx.enrollment_id if kind == "kirim" else None
        if fields["enrollment_id"]:
            fields["program_id"] = tx.program_id
    tx.kind, tx.amount, tx.account_id = kind, amount, account.id
    for k, v in fields.items():
        setattr(tx, k, v)
    # summa yoki hisob o'zgarmasa — yozilgan kundagi kurs saqlanadi (o'tgan oylar siljimasin)
    if money_changed:
        tx.amount_uzs = round(F.to_uzs(amount, account.currency), 2)
    db.session.flush()
    audit.log("tx_update", f"{before} → {audit.describe_tx(tx)}")
    db.session.commit()
    _warn(tx)
    flash("Yozuv yangilandi.", "ok")
    return _back()


def _warn(tx):
    """Saqlangandan keyin ogohlantirish (bloklamaydi): takror, manfiy qoldiq, kelajak sanasi."""
    n = lambda v: f"{abs(v):,.0f}".replace(",", " ")  # noqa: E731
    same = Transaction.query.filter(
        Transaction.id != tx.id, Transaction.date == tx.date, Transaction.kind == tx.kind,
        Transaction.account_id == tx.account_id, func.abs(Transaction.amount - tx.amount) < 0.01,
        func.lower(func.coalesce(Transaction.counterparty, "")) == (tx.counterparty or "").lower(),
    ).count()
    if same:
        flash(f"Diqqat: {tx.date:%d.%m.%Y} kuni xuddi shunday yozuv ({tx.counterparty or tx.kind_name}, "
              f"{n(tx.amount)}) avval ham kiritilgan — takror emasligini tekshiring.", "warn")
    bal = F.account_balances().get(tx.account_id, 0)
    if bal < -0.01:
        flash(f"«{tx.account.name}» qoldig'i manfiy bo'lib qoldi: −{n(bal)}. Summa yoki hisobni tekshiring.", "warn")
    if tx.date > date.today():
        flash(f"Sana kelajakda ({tx.date:%d.%m.%Y}) — shu kungacha hisobotlarda ko'rinmaydi.", "warn")


@bp.route("/excel")
def export():
    if g.user.role == "menejer":
        flash("Eksport uchun ruxsat yo'q.", "bad")
        return redirect(url_for("kassa.ledger"))
    fl = _filters()
    month, kind = fl["month"], fl["kind"]
    rows = _kind_filter(fl["base"], kind).order_by(Transaction.date, Transaction.id).all()
    account = db.session.get(Account, fl["acc_id"]) if fl["acc_id"] else None
    all_rows = fl["period"].order_by(Transaction.date, Transaction.id).all()
    opening, cur, flows = F.ledger_rows(all_rows, account=account, start=fl["start"])
    wb = Workbook()
    ws = wb.active
    ws.title = f"Kassa {month:%m.%Y}"
    unit = "$" if cur == "USD" else "so'm"
    # Sheets'dagi kassa daftari tartibi; qayta import qilsa ham ustunlar taniladi
    head = ["Sana", "Turi", "Hisob", "Kimdan / kimga", "Kategoriya", "Dastur", "Izoh",
            f"Kirim, {unit}", f"Chiqim, {unit}", f"Qoldiq, {unit}", "Valyuta", "Aqd", "Shar'iy holat", "Kiritgan"]
    ws.append(head)
    bold = Font(bold=True)
    for c in ws[1]:
        c.font, c.fill = bold, PatternFill("solid", fgColor="E0F1E9")
    ws.append([fl["start"], "", account.name if account else "Barcha hisoblar", "Oy boshidagi qoldiq",
               "", "", "", None, None, round(opening, 2)])
    for c in ws[2]:
        c.fill = PatternFill("solid", fgColor="FBF3DF")
    for t in rows:
        r = flows[t.id]
        refund = t.kind == "kirim" and t.amount < 0
        ws.append([
            t.date, "Qaytarish" if refund else t.kind_name,
            f"{t.account.name} → {t.to_account.name}" if t.kind == "otkazma" else t.account.name,
            t.counterparty, t.category.name if t.category else "", t.program.name if t.program else "", t.note,
            (-r["out"] if refund else (r["inc"] or None)) if not r["move"] else None,
            None if refund or r["move"] else (r["out"] or None),
            round(r["bal"], 2), t.currency, AQD_TYPES.get(t.aqd, ""), SHARIAT.get(t.shariat, ""),
            t.author.name if t.author else "",
        ])
    last = ws.max_row
    ws.append(["", "", "", f"Jami: {len(rows)} ta yozuv", "", "", "",
               f"=SUM(H3:H{last})" if last >= 3 else 0, f"=SUM(I3:I{last})" if last >= 3 else 0,
               round(flows[all_rows[-1].id]["bal"] if all_rows else opening, 2)])
    for c in ws[ws.max_row]:
        c.font = bold
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD.MM.YYYY"
        for k in (7, 8, 9):
            row[k].number_format = "# ##0" if cur == "UZS" else "#,##0.00"
    for i, w in enumerate((12, 12, 22, 24, 26, 26, 30, 15, 15, 16, 8, 16, 12, 14), start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"assoodiq-kassa-{month:%Y-%m}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/import", methods=["GET", "POST"])
def import_rows():
    """Google Sheets / Excel'dan ko'chirish: qo'yish → tekshirish → saqlash."""
    if g.user.role == "menejer":
        flash("Import uchun ruxsat yo'q.", "bad")
        return redirect(url_for("kassa.ledger"))
    accounts = Account.query.filter_by(active=True).order_by(Account.sort, Account.id).all()
    f = request.form
    text = f.get("text") or ""
    default_acc_id = form_int(f.get("default_account")) or (accounts[0].id if accounts else None)
    default_kind = f.get("default_kind") if f.get("default_kind") in ("kirim", "chiqim") else "kirim"
    analysis = None
    if request.method == "POST":
        upload = request.files.get("file")
        if upload and upload.filename:
            try:
                text = sheets.file_to_text(upload)
            except ValueError as e:
                flash(str(e), "bad")
        analysis = sheets.analyze(
            text, accounts,
            Category.query.filter_by(active=True).order_by(Category.sort).all(),
            Program.query.order_by(Program.id.desc()).all(),
            default_account=db.session.get(Account, default_acc_id or 0), default_kind=default_kind,
        )
        if not analysis["rows"] and text.strip():
            flash("Qatorlarni ajratib bo'lmadi — jadvaldan nusxa olib qo'ying.", "bad")
        if f.get("action") == "save":
            good = [r for r in analysis["rows"] if r["status"] == "ok"]
            if not good:
                flash("Saqlanadigan yangi yozuv yo'q.", "bad")
            else:
                for r in good:
                    F.record(r["kind"], r["amount"], r["account"], date=r["date"],
                             category_id=r["category"].id if r["category"] else None,
                             program_id=r["program"].id if r["program"] else None,
                             counterparty=r["counterparty"], note=r["note"], aqd="", shariat="halol",
                             created_by=g.user.id)
                skipped = len(analysis["rows"]) - len(good)
                audit.log("tx_import", f"{len(good)} ta yozuv qo'shildi, {skipped} ta o'tkazib yuborildi")
                db.session.commit()
                flash(f"{len(good)} ta yozuv kassaga qo'shildi"
                      + (f", {skipped} ta qator o'tkazib yuborildi." if skipped else "."), "ok")
                latest = max(r["date"] for r in good)
                return redirect(url_for("kassa.ledger", oy=latest.strftime("%Y-%m")))
    return render_template(
        "pages/kassa_import.html", text=text, analysis=analysis, accounts=accounts,
        default_acc_id=default_acc_id, default_kind=default_kind,
        default_labels=sheets.DEFAULT_LABELS,
    )


@bp.route("/import/namuna.xlsx")
def import_template():
    """Bo'sh shablon: ustunlar tartibi va ikki namuna qator."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Kassa"
    ws.append(sheets.DEFAULT_LABELS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E0F1E9")
    acc = Account.query.filter_by(active=True).order_by(Account.sort, Account.id).first()
    today = date.today()
    ws.append([today, "Kirim", 1500000, acc.name if acc else "", "Kurs to'lovlari", "Alisher Karimov", "3-potok, 1-qism"])
    ws.append([today, "Chiqim", 450000, acc.name if acc else "", "Reklama (Instagram, Telegram)", "Instagram", "Oktabr reklamasi"])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD.MM.YYYY"
        row[2].number_format = "# ##0"
    for i, w in enumerate((12, 10, 14, 20, 30, 24, 30), start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="assoodiq-kassa-shablon.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/<int:tx_id>/ochirish", methods=["POST"])
def delete(tx_id):
    if g.user.role == "menejer":
        flash("O'chirish uchun ruxsat yo'q.", "bad")
        return _back()
    _delete_tx(db.get_or_404(Transaction, tx_id))
    db.session.commit()
    flash("Yozuv o'chirildi.", "ok")
    return _back()


def _delete_tx(tx):
    debt = tx.debt
    audit.log("tx_delete", audit.describe_tx(tx))
    db.session.delete(tx)
    db.session.flush()
    if debt is not None:  # qaytarish o'chirilsa qarz qayta ochiladi
        debt.closed = F.debt_remaining(debt) <= 0.01


@bp.route("/ommaviy", methods=["POST"])
def bulk():
    """Belgilangan bir nechta yozuv: kategoriyasini qo'yish yoki o'chirish."""
    if g.user.role == "menejer":
        flash("Ruxsat yo'q.", "bad")
        return _back()
    ids = [i for i in (form_int(x) for x in request.form.getlist("ids")) if i]
    txs = Transaction.query.filter(Transaction.id.in_(ids)).all() if ids else []
    if not txs:
        flash("Hech qaysi yozuv tanlanmagan.", "bad")
        return _back()
    action = request.form.get("action")
    if action == "category":
        cat = db.session.get(Category, form_int(request.form.get("category_id")) or 0)
        if cat is None:
            flash("Kategoriyani tanlang.", "bad")
            return _back()
        ok = [t for t in txs if t.kind == cat.kind]
        for t in ok:
            t.category_id = cat.id
        skipped = len(txs) - len(ok)
        audit.log("tx_update", f"{len(ok)} ta yozuv kategoriyasi → «{cat.name}»")
        db.session.commit()
        flash(f"{len(ok)} ta yozuvga «{cat.name}» qo'yildi"
              + (f"; {skipped} tasi turi mos kelmagani uchun o'zgarmadi." if skipped else "."), "ok" if ok else "warn")
    elif action == "delete":
        for t in txs:
            _delete_tx(t)
        db.session.commit()
        flash(f"{len(txs)} ta yozuv o'chirildi.", "ok")
    return _back()


def _back_url():
    return safe_next(url_for("kassa.ledger"))


def _back():
    return redirect(_back_url())

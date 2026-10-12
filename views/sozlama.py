import secrets
from datetime import date
from io import BytesIO

from flask import Blueprint, flash, g, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from core import audit
from core import finance as F
from core.extensions import db
from core.models import (
    ACCOUNT_KINDS, AQD_TYPES, PROGRAM_KINDS, ROLES, SHARIAT, Account, AuditLog, Category, Debt,
    Enrollment, Member, Program, Transaction, User,
)
from views import form_int, money, require_role

bp = Blueprint("sozlama", __name__, url_prefix="/sozlamalar")

SETTING_KEYS = ("org_name", "usd_rate", "gold_price_g", "silver_price_g", "nisob_basis",
                "inventory_value", "sadaqa_percent")
NUMERIC = {"usd_rate", "gold_price_g", "silver_price_g", "inventory_value", "sadaqa_percent"}


@bp.route("/")
def index():
    return render_template(
        "pages/settings.html",
        s={k: F.get_setting(k) for k in SETTING_KEYS},
        accounts=F.accounts_with_balance(active_only=False),
        categories=Category.query.order_by(Category.kind, Category.sort, Category.name).all(),
        users=User.query.order_by(User.id).all(), roles=ROLES, account_kinds=ACCOUNT_KINDS,
        program_kinds=PROGRAM_KINDS, tab=request.args.get("bolim", "umumiy"),
        stats=_data_stats(),
        logs=(AuditLog.query.order_by(AuditLog.created_at.desc()).limit(200).all()
              if request.args.get("bolim") == "tarix" else []),
        actions=audit.ACTIONS,
        tg=_telegram_context() if request.args.get("bolim") == "telegram" else None,
        app=_app_context() if request.args.get("bolim") == "ilova" else None,
        aic=_ai_context() if request.args.get("bolim") == "ai" else None,
    )


def _ai_context():
    import os

    from core import ai
    from core import ai_report as R
    from core.models import StaffLink, User
    from views.ai import usage_stats
    key = F.get_setting("ai_key")
    links = {lk.user_id: lk for lk in StaffLink.query.filter_by(active=True).all()}
    return {"env": bool(os.environ.get("ANTHROPIC_API_KEY")), "on": ai.enabled(),
            "mask": (key[:10] + "…" + key[-4:]) if len(key) > 16 else "", "model": ai.MODEL,
            "opts": {k: F.get_setting(k) for k in ("ai_weekly", "ai_weekly_day", "ai_weekly_hour", "ai_daily_limit",
                                                    "contract_students", "contract_student_daily")},
            "usage": usage_stats(), "staff": User.query.filter_by(active=True).order_by(User.id).all(),
            "links": links, "recipients": len(R.recipients())}


def _app_context():
    from datetime import date

    from core import bunny
    from core import prayer as P
    from core.models import StudentAccess
    from views.oquv import app_url

    def mask(v):
        return (v[:4] + "…" + v[-4:]) if len(v) > 10 else ("•" * len(v))
    ps = P.settings()
    lat, lng = P.CITIES[ps["city"]]
    return {"bunny": bunny.conf(), "bunny_on": bunny.enabled(), "mask": mask, "ps": ps,
            "methods": P.METHODS, "cities": P.CITIES, "prayers": P.PRAYERS, "link": app_url(),
            "preview": P.day_schedule(date.today(), lat, lng, method=ps["method"], asr=ps["asr"], offsets=ps["offsets"]),
            "users": StudentAccess.query.filter(StudentAccess.last_seen.isnot(None)).count()}


def _telegram_context():
    import os
    from datetime import datetime, timezone
    from core import telegram as T
    from core.models import Member, Notification, TelegramLink
    tok = T.token()
    return {
        "env": bool(os.environ.get("TELEGRAM_BOT_TOKEN")),
        "token_mask": (tok[:6] + "…" + tok[-4:]) if tok else "",
        "bot": F.get_setting("tg_bot"), "webhook": F.get_setting("tg_webhook"),
        "opts": {k: F.get_setting(k) for k in ("tg_auto_lessons", "tg_lead_hours", "tg_auto_expiry",
                                               "tg_expiry_days", "tg_auto_pay", "tg_pay_day")},
        "tpl": {k: T.template(k) for k in T.DEFAULT_TEMPLATES}, "ph": T.PLACEHOLDERS,
        "linked": TelegramLink.query.filter_by(active=True).count(), "members": Member.query.count(),
        "sent_today": Notification.query.filter(Notification.status == "sent",
                                                Notification.created_at >= datetime.now(timezone.utc).replace(tzinfo=None) - __import__("datetime").timedelta(hours=24)).count(),
        "log": Notification.query.order_by(Notification.created_at.desc()).limit(25).all(),
        "site": request.url_root.rstrip("/"),
    }


@bp.route("/umumiy", methods=["POST"])
def save_general():
    changed = []
    for k in SETTING_KEYS:
        if k in request.form:
            v = request.form[k].strip()
            if k in NUMERIC:
                v = str(money(v))
            if k == "nisob_basis" and v not in F.NISOB_GRAMS:
                v = "oltin"
            if F.get_setting(k) != v:
                changed.append(f"{k}: {F.get_setting(k)} → {v}")
            F.set_setting(k, v)
    if changed:
        audit.log("settings", "; ".join(changed))
    db.session.commit()
    flash("Sozlamalar saqlandi.", "ok")
    return redirect(url_for("sozlama.index"))


@bp.route("/namuna", methods=["POST"])
def demo_load():
    """Bo'sh bazaga namuna (test) ma'lumot yuklash — faqat rahbar."""
    require_role("rahbar")
    from core.demo import load_demo
    if load_demo():
        audit.log("demo_load", "8 dastur, 46 a'zo, 12 oylik yozuvlar")
        db.session.commit()
        flash("Namuna ma'lumotlar yuklandi: 8 dastur, 46 a'zo, 12 oylik kirim-chiqim. Haqiqiy ishdan oldin tozalang.", "ok")
    else:
        flash("Bazada allaqachon yozuvlar bor — namuna faqat bo'sh bazaga yuklanadi.", "bad")
    return redirect(url_for("main.dashboard"))


@bp.route("/tozalash", methods=["POST"])
def data_clear():
    """Barcha moliyaviy yozuvlarni o'chirish (haqiqiy ishni toza boshlash uchun)."""
    require_role("rahbar")
    if (request.form.get("tasdiq") or "").strip().upper() != "TOZALASH":
        flash("Tasdiqlash uchun TOZALASH so'zini yozing.", "bad")
        return redirect(url_for("sozlama.index", bolim="malumotlar"))
    from core.demo import clear_data
    clear_data()
    audit.log("data_clear", "barcha moliyaviy yozuvlar o'chirildi")
    db.session.commit()
    flash("Barcha yozuvlar o'chirildi. Xodimlar, hisoblar va kategoriyalar saqlandi — haqiqiy ishni boshlashingiz mumkin.", "ok")
    return redirect(url_for("sozlama.index", bolim="malumotlar"))


@bp.route("/zaxira")
def backup():
    """Barcha ma'lumotlarning Excel nusxasi (zaxira yoki boshqa tizimga ko'chirish uchun)."""
    wb = Workbook()
    head_font, head_fill = Font(bold=True), PatternFill("solid", fgColor="E0F1E9")

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        for i in range(1, len(header) + 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = 18

    sheet("Kassa", ["ID", "Sana", "Turi", "Hisob", "Qaysi hisobga", "Summa", "Valyuta", "So'mda", "Kategoriya",
                    "Dastur", "Kimdan/kimga", "Aqd", "Shar'iy holat", "Izoh", "Kiritgan", "Kiritilgan vaqt"],
          ([t.id, t.date, t.kind_name, t.account.name, t.to_account.name if t.to_account else "", t.amount,
            t.currency, t.amount_uzs, t.category.name if t.category else "", t.program.name if t.program else "",
            t.counterparty, AQD_TYPES.get(t.aqd, ""), SHARIAT.get(t.shariat, ""), t.note,
            t.author.name if t.author else "", t.created_at]
           for t in Transaction.query.order_by(Transaction.date, Transaction.id)), first=True)
    sheet("Dasturlar", ["ID", "Nomi", "Turi", "To'lov tartibi", "Narx", "Valyuta", "Boshlanishi", "Tugashi",
                        "O'rinlar", "Holati"],
          ([p.id, p.name, p.kind_name, p.billing_name, p.price, p.currency, p.start_date, p.end_date,
            p.capacity, p.status_name] for p in Program.query.order_by(Program.id)))
    sheet("A'zolar", ["ID", "Ism-familiya", "Telefon", "Biznesi", "Shahar", "Izoh"],
          ([m.id, m.full_name, m.phone, m.business, m.city, m.note] for m in Member.query.order_by(Member.id)))
    sheet("Yozilishlar", ["A'zo", "Dastur", "Hisoblangan narx", "Qo'shilgan", "Obuna tugashi", "Holati"],
          ([e.member.full_name, e.program.name, e.price, e.joined_at, e.valid_until, e.status]
           for e in Enrollment.query.order_by(Enrollment.id)))
    sheet("Qarzlar", ["ID", "Turi", "Kim", "Telefon", "Summa", "Valyuta", "Qoldiq", "Sana", "Muddat", "Guvoh",
                      "Izoh", "Yopilgan"],
          ([d.id, d.direction_name, d.counterparty, d.phone, d.amount, d.currency, F.debt_remaining(d), d.date,
            d.due_date, d.witness, d.note, "ha" if d.closed else "yo'q"] for d in Debt.query.order_by(Debt.id)))
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, date):
                    c.number_format = "DD.MM.YYYY"
                elif isinstance(c.value, float):
                    c.number_format = "# ##0.##"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"assoodiq-zaxira-{date.today():%Y-%m-%d}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/hisob", methods=["POST"])
def account_add():
    f = request.form
    name = (f.get("name") or "").strip()
    if not name:
        flash("Hisob nomini yozing.", "bad")
        return redirect(url_for("sozlama.index", bolim="hisoblar"))
    db.session.add(Account(name=name[:120], kind=f.get("kind") if f.get("kind") in ACCOUNT_KINDS else "boshqa",
                           currency="USD" if f.get("currency") == "USD" else "UZS",
                           opening_balance=money(f.get("opening_balance")),
                           sort=Account.query.count()))
    db.session.commit()
    flash("Hisob qo'shildi.", "ok")
    return redirect(url_for("sozlama.index", bolim="hisoblar"))


@bp.route("/hisob/<int:aid>/holat", methods=["POST"])
def account_toggle(aid):
    a = db.get_or_404(Account, aid)
    a.active = not a.active
    db.session.commit()
    return redirect(url_for("sozlama.index", bolim="hisoblar"))


@bp.route("/kategoriya", methods=["POST"])
def category_add():
    f = request.form
    name = (f.get("name") or "").strip()
    if not name:
        flash("Kategoriya nomini yozing.", "bad")
        return redirect(url_for("sozlama.index", bolim="kategoriyalar"))
    kind = "chiqim" if f.get("kind") == "chiqim" else "kirim"
    db.session.add(Category(name=name[:120], kind=kind, group=(f.get("group") or "Boshqa").strip()[:60],
                            program_kind=f.get("program_kind") if f.get("program_kind") in PROGRAM_KINDS else None))
    db.session.commit()
    flash("Kategoriya qo'shildi.", "ok")
    return redirect(url_for("sozlama.index", bolim="kategoriyalar"))


@bp.route("/kategoriya/<int:cid>/holat", methods=["POST"])
def category_toggle(cid):
    c = db.get_or_404(Category, cid)
    if c.system_code:
        flash("Tizim kategoriyasini o'chirib bo'lmaydi (zakot, tozalash, sadaqa hisobi unga bog'liq).", "bad")
    else:
        c.active = not c.active
        db.session.commit()
    return redirect(url_for("sozlama.index", bolim="kategoriyalar"))


@bp.route("/xodim", methods=["POST"])
def user_add():
    require_role("rahbar")
    f = request.form
    name = (f.get("name") or "").strip()
    if not name:
        flash("Xodim ismini yozing.", "bad")
        return redirect(url_for("sozlama.index", bolim="xodimlar"))
    code = _new_code()
    u = User(name=name[:120], role=f.get("role") if f.get("role") in ROLES else "menejer")
    u.set_code(code)
    db.session.add(u)
    audit.log("user_add", f"{u.name} ({u.role_name})")
    db.session.commit()
    flash(f"{u.name} qo'shildi. Kirish kodi: {code} — uni xodimga bering, keyin ko'rinmaydi.", "ok")
    return redirect(url_for("sozlama.index", bolim="xodimlar"))


@bp.route("/xodim/<int:uid>/kod", methods=["POST"])
def user_reset(uid):
    require_role("rahbar")
    u = db.get_or_404(User, uid)
    code = _new_code()
    u.set_code(code)
    audit.log("user_code", u.name)
    db.session.commit()
    flash(f"{u.name} uchun yangi kod: {code}", "ok")
    return redirect(url_for("sozlama.index", bolim="xodimlar"))


@bp.route("/xodim/<int:uid>/holat", methods=["POST"])
def user_toggle(uid):
    require_role("rahbar")
    u = db.get_or_404(User, uid)
    if u.id == g.user.id:
        flash("O'zingizni faolsizlantira olmaysiz.", "bad")
    else:
        u.active = not u.active
        audit.log("user_toggle", f"{u.name}: {'faol' if u.active else 'faolsiz'}")
        db.session.commit()
    return redirect(url_for("sozlama.index", bolim="xodimlar"))


def _data_stats():
    from core.models import Debt, Member, Program, Transaction
    return {"tx": Transaction.query.count(), "programs": Program.query.count(),
            "members": Member.query.count(), "debts": Debt.query.count()}


def _new_code():
    """Boshqa xodimnikiga to'g'ri kelmaydigan 6 xonali kod."""
    users = User.query.all()
    while True:
        code = f"{secrets.randbelow(900000) + 100000}"
        if not any(u.check_code(code) for u in users):
            return code

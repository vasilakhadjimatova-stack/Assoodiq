"""Google Sheets / Excel'dan yozuvlarni ko'chirish.

Foydalanuvchi jadvaldan qatorlarni nusxa olib (Ctrl+C) qo'yadi — matn tab bilan
ajratilgan bo'ladi. Sarlavha qatori bo'lsa ustunlar nomi bo'yicha topiladi,
bo'lmasa standart tartib olinadi. Har bir qator tekshiriladi: yaroqli, takror
(bazada allaqachon bor) yoki xato (sababi bilan).
"""
import csv
import io
import re
from collections import Counter
from datetime import date, datetime

from core.models import Transaction

MAX_ROWS = 2000

# Sarlavhadagi so'z → maydon. Tartib muhim: aniqrog'i oldin.
HEADER_WORDS = [
    ("sana", "date"), ("date", "date"), ("kun", "date"),
    ("so'mda", "skip"), ("somda", "skip"), ("valyuta", "skip"), ("qoldiq", "skip"), ("balans", "skip"),
    ("kiritgan", "skip"), ("aqd", "skip"), ("shar", "skip"), ("№", "skip"),
    ("turi", "kind"), ("tur", "kind"), ("type", "kind"),
    ("kirim", "income"), ("tushum", "income"), ("income", "income"),
    ("chiqim", "expense"), ("xarajat", "expense"), ("expense", "expense"),
    ("summa", "amount"), ("miqdor", "amount"), ("amount", "amount"),
    ("hisob", "account"), ("kassa", "account"), ("account", "account"),
    ("kategoriya", "category"), ("modda", "category"), ("category", "category"),
    ("dastur", "program"), ("program", "program"),
    ("kimdan", "counterparty"), ("kimga", "counterparty"), ("kim", "counterparty"),
    ("mijoz", "counterparty"), ("kontragent", "counterparty"), ("ism", "counterparty"),
    ("izoh", "note"), ("tavsif", "note"), ("sharh", "note"), ("note", "note"), ("comment", "note"),
]
# Sarlavhasiz qo'yilganda ustunlar tartibi (sahifada ham shunday ko'rsatiladi)
DEFAULT_ORDER = ["date", "kind", "amount", "account", "category", "counterparty", "note"]
DEFAULT_LABELS = ["Sana", "Turi", "Summa", "Hisob", "Kategoriya", "Kimdan / kimga", "Izoh"]


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip().lower().replace("‘", "'").replace("’", "'").replace("ʻ", "'").replace("`", "'"))


def _loose(s):
    """Taqqoslash uchun: apostrof va bo'shliqlarsiz."""
    return re.sub(r"[\s'\"().,-]", "", _norm(s))


def text_to_rows(text):
    """Qo'yilgan matn → qatorlar ro'yxati (tab, nuqtali vergul yoki vergul bilan)."""
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip("\n")
    if not text.strip():
        return []
    sample = text[:4000]
    if "\t" in sample:
        delim = "\t"
    elif ";" in sample:
        delim = ";"
    else:
        delim = ","
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    rows = [[c.strip() for c in r] for r in reader]
    return [r for r in rows if any(c for c in r)]


def file_to_text(storage):
    """Yuklangan .xlsx / .csv fayl → tab bilan ajratilgan matn (sahifadagi maydonga qo'yiladi)."""
    name = (storage.filename or "").lower()
    data = storage.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024:
        raise ValueError("Fayl juda katta (5 MB gacha).")
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 — buzilgan fayl
            raise ValueError("Excel faylni o'qib bo'lmadi.")
        ws = wb.worksheets[0]
        lines = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_text(v) for v in row]
            if any(cells):
                lines.append("\t".join(c.replace("\t", " ").replace("\n", " ") for c in cells))
            if len(lines) > MAX_ROWS + 1:
                break
        return "\n".join(lines)
    if name.endswith((".csv", ".tsv", ".txt")):
        for enc in ("utf-8-sig", "cp1251"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        raise ValueError("Fayl kodirovkasini aniqlab bo'lmadi (UTF-8 da saqlang).")
    raise ValueError("Faqat .xlsx yoki .csv fayl qabul qilinadi.")


def _cell_text(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float):
        return f"{v:.2f}".rstrip("0").rstrip(".")
    return str(v).strip()


def detect_header(row):
    """Sarlavha qatori bo'lsa {ustun_raqami: maydon}, aks holda None."""
    if any(parse_date(c) or re.fullmatch(r"[-+−(]?[\d\s.,]+\)?", str(c).strip()) for c in row if str(c).strip()):
        return None  # sana yoki son bor — bu ma'lumot qatori
    mapping, used = {}, set()
    for i, cell in enumerate(row):
        n = _norm(cell)
        if not n or re.search(r"\d", n.replace("№", "")):
            continue
        for word, field in HEADER_WORDS:
            if word in n:
                if field != "skip" and field not in used:
                    mapping[i] = field
                    used.add(field)
                break
    has_money = used & {"amount", "income", "expense"}
    return mapping if has_money and len(used) >= 2 else None


def parse_amount(raw):
    """'1 500 000', '1,500,000', '1.500.000', '-250 000 so'm', '(300)' → float yoki None."""
    s = str(raw or "").strip()
    if not s:
        return None
    neg = s.startswith(("-", "−", "–")) or (s.startswith("(") and s.endswith(")"))
    s = re.sub(r"(?i)so['‘’ʻ]?m|sum|uzs|usd|\$|сум", "", s)
    s = re.sub(r"[^\d.,]", "", s)
    if not s:
        return None
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", "") if re.fullmatch(r"\d{1,3}(,\d{3})+", s) else s.replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3}){2,}", s) or re.fullmatch(r"\d{1,3}\.\d{3}", s):
        s = s.replace(".", "")
    try:
        v = round(float(s), 2)
    except ValueError:
        return None
    return -v if neg else v


def parse_date(raw, today=None):
    s = str(raw or "").strip()
    if not s:
        return None
    s = s.split(" ")[0] if re.match(r"^\d{4}-\d{2}-\d{2}[ T]", s) else s
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        y, mo, d = (int(x) for x in m.groups())
    else:
        m = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", s)
        if not m:
            return None
        d, mo, y = (int(x) for x in m.groups())
        if mo > 12 and d <= 12:  # oy/kun/yil (AQSh tartibi)
            d, mo = mo, d
        if y < 100:
            y += 2000
    try:
        return date(y, mo, d)
    except ValueError:
        return None


def parse_kind(raw):
    n = _loose(raw)
    if not n:
        return None
    if "otkazma" in n or "transfer" in n or "qarz" in n:
        return "skip"
    if n in ("+",) or "kirim" in n or "tushum" in n or "income" in n:
        return "kirim"
    if n in ("-",) or "chiqim" in n or "xarajat" in n or "expense" in n:
        return "chiqim"
    return None


def _match(name, items, key=lambda x: x.name):
    """Nomi bo'yicha topish: aniq mos, keyin boshlanishi, keyin ichida bo'lishi."""
    n = _loose(name)
    if not n:
        return None
    exact = [x for x in items if _loose(key(x)) == n]
    if exact:
        return exact[0]
    starts = [x for x in items if _loose(key(x)).startswith(n) or n.startswith(_loose(key(x)))]
    if len(starts) == 1:
        return starts[0]
    inside = [x for x in items if n in _loose(key(x))]
    return inside[0] if len(inside) == 1 else None


def analyze(text, accounts, categories, programs, default_account=None, default_kind="kirim"):
    """Matnni tahlil qiladi. Natija: {'rows': [...], 'header': bool, 'columns': [...], 'counts': {...}}."""
    raw_rows = text_to_rows(text)
    result = {"rows": [], "header": False, "columns": [], "counts": Counter(), "too_many": False}
    if not raw_rows:
        return result
    mapping = detect_header(raw_rows[0])
    if mapping:
        result["header"] = True
        raw_rows = raw_rows[1:]
    else:
        mapping = dict(enumerate(DEFAULT_ORDER))
    labels = dict(zip(DEFAULT_ORDER, DEFAULT_LABELS))
    labels.update(income="Kirim", expense="Chiqim", program="Dastur")
    result["columns"] = [labels.get(f, f) for _, f in sorted(mapping.items())]
    if len(raw_rows) > MAX_ROWS:
        result["too_many"] = True
        raw_rows = raw_rows[:MAX_ROWS]

    cats_by_kind = {k: [c for c in categories if c.kind == k] for k in ("kirim", "chiqim")}
    fallback_cat = {k: next((c for c in cats_by_kind[k] if c.name.lower().startswith("boshqa")), None)
                    for k in ("kirim", "chiqim")}

    for n, cells in enumerate(raw_rows, start=2 if result["header"] else 1):
        get = {f: (cells[i] if i < len(cells) else "") for i, f in mapping.items()}
        row = {"n": n, "cells": cells, "status": "ok", "msg": [], "warn": []}
        result["rows"].append(row)

        d = parse_date(get.get("date"))
        if d is None:
            row["status"], row["msg"] = "err", ["sana tushunarsiz" if get.get("date") else "sana yo'q"]
            continue

        kind = parse_kind(get.get("kind"))
        amount = None
        if kind == "skip":
            row["status"], row["msg"] = "skip", ["o'tkazma/qarz — qo'lda yozing"]
            continue
        if "income" in get or "expense" in get:
            inc, exp = parse_amount(get.get("income")), parse_amount(get.get("expense"))
            inc, exp = abs(inc or 0), abs(exp or 0)
            if inc and exp:
                row["status"], row["msg"] = "err", ["kirim va chiqim ikkalasi to'ldirilgan"]
                continue
            if inc or exp:
                amount = inc or exp
                kind = kind or ("kirim" if inc else "chiqim")
        if amount is None:
            v = parse_amount(get.get("amount"))
            if v is not None:
                amount = abs(v)
                if kind is None:
                    kind = "chiqim" if v < 0 else default_kind
        if not amount:
            row["status"], row["msg"] = "err", ["summa yo'q" if not any(get.get(k) for k in ("amount", "income", "expense")) else "summa tushunarsiz"]
            continue
        kind = kind or default_kind

        account = default_account
        if get.get("account"):
            account = _match(get["account"], accounts)
            if account is None:
                row["status"], row["msg"] = "err", [f"«{get['account']}» hisobi topilmadi"]
                continue
        if account is None:
            row["status"], row["msg"] = "err", ["hisob ko'rsatilmagan"]
            continue

        category = None
        if get.get("category"):
            category = _match(get["category"], cats_by_kind[kind])
            if category is None:
                category = fallback_cat[kind]
                row["warn"].append(f"«{get['category']}» kategoriyasi yo'q → {category.name if category else 'kategoriyasiz'}")
        else:
            category = fallback_cat[kind]

        program = _match(get.get("program"), programs) if get.get("program") else None
        if get.get("program") and program is None:
            row["warn"].append(f"«{get['program']}» dasturi topilmadi")

        row.update(date=d, kind=kind, amount=amount, account=account, category=category, program=program,
                   counterparty=(get.get("counterparty") or "")[:160], note=(get.get("note") or "")[:500])

    _mark_duplicates(result["rows"])
    for r in result["rows"]:
        result["counts"][r["status"]] += 1
    return result


def _mark_duplicates(rows):
    """Bazada xuddi shunday yozuv (sana, tur, hisob, summa, kim) bo'lsa — takror, saqlanmaydi.

    Bir xil to'lov ikki marta bo'lishi mumkin, shuning uchun bazadagi soniga qarab belgilanadi.
    """
    good = [r for r in rows if r["status"] == "ok"]
    if not good:
        return
    lo, hi = min(r["date"] for r in good), max(r["date"] for r in good)
    existing = Counter(
        (t.date, t.kind, t.account_id, round(t.amount, 2), _loose(t.counterparty))
        for t in Transaction.query.filter(Transaction.date >= lo, Transaction.date <= hi,
                                          Transaction.kind.in_(("kirim", "chiqim"))).all()
    )
    for r in good:
        key = (r["date"], r["kind"], r["account"].id, round(r["amount"], 2), _loose(r["counterparty"]))
        if existing[key] > 0:
            existing[key] -= 1
            r["status"] = "dup"
            r["msg"] = ["bazada bor — o'tkazib yuboriladi"]

from datetime import date, timedelta

import pytest

from app import create_app
from config import TestConfig
from core import finance as F
from core import hijri
from core.extensions import db
from core.models import Account, Category, Debt, Enrollment, Member, Program, Transaction, User
from core.bootstrap import base_setup
from seed import DEMO_USERS


@pytest.fixture
def app():
    app = create_app(TestConfig)
    with app.app_context():
        base_setup(users=DEMO_USERS)
        yield app
        db.session.remove()
        db.drop_all()


def login(client, code):
    return client.post("/kirish", data={"code": code})


def acc(name):
    return Account.query.filter_by(name=name).one()


def cat(name):
    return Category.query.filter_by(name=name).one()


# ── Hijriy taqvim ──

def test_hijri_roundtrip():
    d = date(2024, 1, 1)
    for _ in range(800):
        assert hijri.from_hijri(*hijri.to_hijri(d)) == d
        d += timedelta(days=3)


def test_hijri_known_date():
    # 1 Ramazon 1445 ≈ 11 mart 2024 (jadval usuli ±1 kun)
    assert abs((hijri.from_hijri(1445, 9, 1) - date(2024, 3, 11)).days) <= 1


def test_next_hawl_is_one_lunar_year():
    start = date(2025, 3, 1)
    nxt = hijri.next_hawl(start, today=start + timedelta(days=1))
    assert 353 <= (nxt - start).days <= 356


# ── Zakot ──

def test_zakat_below_nisob_is_zero():
    r = F.zakat_compute(cash=50_000_000, receivables=0, inventory=0, other_assets=0,
                        liabilities=0, nisob=100_000_000)
    assert r["due"] == 0 and not r["reached"]


def test_zakat_deducts_liabilities_and_takes_quarter_tenth():
    r = F.zakat_compute(cash=200_000_000, receivables=40_000_000, inventory=10_000_000,
                        other_assets=0, liabilities=50_000_000, nisob=100_000_000)
    assert r["base"] == 200_000_000
    assert r["due"] == 5_000_000


def test_zakat_defaults_exclude_purification_fund(app):
    F.record("kirim", 1_000_000, acc("Asosiy kassa"), category_id=cat("Boshqa kirim").id, shariat="nohalol")
    db.session.commit()
    d = F.zakat_defaults()
    assert d["cash"] == F.total_balance_uzs() - 1_000_000


# ── Daftar va balans ──

def test_balances_and_pnl_ignore_transfers_and_debts(app):
    kassa, bank, usd = acc("Asosiy kassa"), acc("Bank hisobi"), acc("Dollar kassa")
    start_kassa = F.account_balances()[kassa.id]
    F.record("kirim", 5_000_000, kassa, category_id=cat("Kurs to'lovlari").id)
    F.record("chiqim", 1_000_000, kassa, category_id=cat("Ofis ijarasi").id)
    F.record("otkazma", 2_000_000, kassa, to_account_id=bank.id, to_amount=2_000_000)
    F.record("qarz_chiqim", 500_000, kassa)
    F.record("kirim", 100, usd, category_id=cat("Konsalting").id)
    db.session.commit()

    bal = F.account_balances()
    assert bal[kassa.id] == start_kassa + 5_000_000 - 1_000_000 - 2_000_000 - 500_000
    t = F.period_totals(date.today() - timedelta(days=1), date.today())
    assert t["income"] == 5_000_000 + 100 * F.usd_rate()
    assert t["expense"] == 1_000_000


def test_amount_uzs_frozen_at_record_rate(app):
    F.set_setting("usd_rate", "12000")
    tx = F.record("kirim", 100, acc("Dollar kassa"), category_id=cat("Konsalting").id)
    db.session.commit()
    F.set_setting("usd_rate", "13000")
    db.session.commit()
    assert tx.amount_uzs == 1_200_000


def test_parse_money_handles_spaces_and_commas():
    assert F.parse_money("1 500 000") == 1_500_000
    assert F.parse_money("1 250,5") == 1250.5
    assert F.parse_money("abc") == 0


# ── Dasturlar ──

def test_enrollment_due_and_program_profit(app):
    p = Program(name="Kurs", kind="kurs", price=4_000_000)
    m = Member(full_name="Test A'zo")
    db.session.add_all([p, m])
    db.session.flush()
    e = Enrollment(member_id=m.id, program_id=p.id, price=4_000_000)
    db.session.add(e)
    db.session.flush()
    F.record("kirim", 1_500_000, acc("Payme"), enrollment_id=e.id, program_id=p.id,
             category_id=cat("Kurs to'lovlari").id)
    F.record("chiqim", 600_000, acc("Korporativ karta"), program_id=p.id,
             category_id=cat("Reklama (Instagram, Telegram)").id)
    db.session.commit()
    assert F.enrollment_row(e)["due"] == 2_500_000
    s = F.program_stats(p)
    assert s["profit"] == 900_000
    assert s["cac"] == 600_000


def test_club_renew_extends_period_and_bills(app):
    client = app.test_client()
    login(client, "100100")
    p = Program(name="Klub", kind="klub", billing="oylik", price=1_000_000)
    m = Member(full_name="Klub a'zosi")
    db.session.add_all([p, m])
    db.session.commit()
    client.post(f"/dasturlar/{p.id}/yozish", data={"member_id": m.id, "joined_at": date.today().isoformat()})
    e = Enrollment.query.one()
    first_until = e.valid_until
    client.post(f"/yozilish/{e.id}/uzaytirish")
    db.session.refresh(e)
    assert e.price == 2_000_000
    assert 27 <= (e.valid_until - first_until).days <= 31


# ── Qarz daftari ──

def test_debt_repayment_closes_debt(app):
    client = app.test_client()
    login(client, "100100")
    bank = acc("Bank hisobi")
    client.post("/qarz/yangi", data={"direction": "berilgan", "counterparty": "Hamkor",
                                     "amount": "3 000 000", "account_id": bank.id})
    d = Debt.query.one()
    assert F.debt_remaining(d) == 3_000_000
    client.post(f"/qarz/{d.id}/qaytarish", data={"amount": "1 000 000", "account_id": bank.id})
    assert not db.session.get(Debt, d.id).closed
    client.post(f"/qarz/{d.id}/qaytarish", data={"amount": "2 000 000", "account_id": bank.id})
    d = db.session.get(Debt, d.id)
    assert d.closed and F.debt_remaining(d) == 0
    # qarz foyda-zararga kirmaydi
    assert F.period_totals(date.today(), date.today())["income"] == 0


# ── Halollik ──

def test_purification_fund(app):
    F.record("kirim", 400_000, acc("Bank hisobi"), category_id=cat("Boshqa kirim").id, shariat="nohalol")
    F.record("chiqim", 150_000, acc("Bank hisobi"), category_id=F.system_category("tozalash").id)
    db.session.commit()
    assert F.purification_fund() == {"flagged": 400_000, "paid": 150_000, "remaining": 250_000}


# ── Kirish va rollar ──

def test_login_required_and_wrong_code(app):
    client = app.test_client()
    assert client.get("/").status_code == 302
    r = login(client, "999999")
    assert "noto" in r.get_data(as_text=True)


def test_login_rate_limited(app):
    from views import auth
    auth._fails.clear()
    client = app.test_client()
    for _ in range(auth.MAX_FAILS):
        login(client, "000000")
    r = login(client, "100100")  # to'g'ri kod ham bloklanadi
    assert r.status_code == 200 and "Juda ko" in r.get_data(as_text=True)
    auth._fails.clear()


def test_manager_permissions(app):
    client = app.test_client()
    login(client, "300300")
    assert client.get("/zakot").status_code == 302
    assert client.get("/dasturlar").status_code == 200
    r = client.post("/kassa/yangi", data={"kind": "chiqim", "amount": "1000", "account_id": acc("Asosiy kassa").id})
    assert r.status_code == 302
    assert Transaction.query.count() == 0


def test_all_pages_render_with_demo_data(app):
    from seed import demo
    demo()
    client = app.test_client()
    login(client, "100100")
    pid, mid, did = Program.query.first().id, Member.query.first().id, Debt.query.first().id
    for url in ["/", "/kassa/", "/kassa/?tur=qarz", "/dasturlar", f"/dasturlar/{pid}", "/azolar",
                f"/azolar/{mid}", "/qarz/", f"/qarz/{did}", "/hisobot/",
                "/hisobot/excel", "/sozlamalar/", "/sozlamalar/?bolim=xodimlar"]:
        assert client.get(url).status_code == 200, url
    for url in ("/zakot", "/halollik"):  # bo'limlar menyudan olib tashlangan
        assert client.get(url).status_code == 302, url


def test_new_user_code_shown_once_and_works(app):
    client = app.test_client()
    login(client, "100100")
    r = client.post("/sozlamalar/xodim", data={"name": "Yangi menejer", "role": "menejer"}, follow_redirects=True)
    text = r.get_data(as_text=True)
    import re
    code = re.search(r"Kirish kodi: (\d{6})", text).group(1)
    u = User.query.filter_by(name="Yangi menejer").one()
    assert u.check_code(code)


def test_first_run_creates_single_admin_from_env(monkeypatch):
    monkeypatch.setenv("ADMIN_CODE", "482915")
    app = create_app(TestConfig)
    with app.app_context():
        base_setup()
        users = User.query.all()
        assert [(u.role, u.check_code("482915")) for u in users] == [("rahbar", True)]
        assert Account.query.count() and Category.query.filter_by(system_code="zakot").count() == 1
        base_setup()  # qayta chaqirish hech narsani takrorlamaydi
        assert User.query.count() == 1
        db.drop_all()


def test_secret_key_persisted_in_db():
    from core.bootstrap import secret_key
    app = create_app(TestConfig)
    with app.app_context():
        k1 = secret_key()
        assert len(k1) == 64 and secret_key() == k1
        db.drop_all()


# ── Test rejimi: namuna yuklash va tozalash ──

def test_demo_load_and_clear(app):
    client = app.test_client()
    login(client, "100100")
    client.post("/sozlamalar/namuna")
    assert Transaction.query.count() > 100 and Program.query.count() == 8
    assert F.get_setting("demo_loaded") == "1"
    assert "Test rejimi" in client.get("/").get_data(as_text=True)

    before = Transaction.query.count()
    r = client.post("/sozlamalar/namuna", follow_redirects=True)  # ikkinchi marta yuklanmaydi
    assert Transaction.query.count() == before and "allaqachon" in r.get_data(as_text=True)

    client.post("/sozlamalar/tozalash", data={"tasdiq": "xato"})  # tasdiqsiz o'chmaydi
    assert Transaction.query.count() == before

    users, cats = User.query.count(), Category.query.count()
    client.post("/sozlamalar/tozalash", data={"tasdiq": "TOZALASH"})
    assert Transaction.query.count() == Program.query.count() == Member.query.count() == Debt.query.count() == 0
    assert User.query.count() == users and Category.query.count() == cats
    assert F.total_balance_uzs() == 0
    assert "Test rejimi" not in client.get("/").get_data(as_text=True)


def test_demo_controls_only_for_rahbar(app):
    client = app.test_client()
    login(client, "200200")  # buxgalter
    assert client.post("/sozlamalar/namuna").status_code == 403
    assert client.post("/sozlamalar/tozalash", data={"tasdiq": "TOZALASH"}).status_code == 403
    assert Transaction.query.count() == 0


# ── Audit: tuzatishlar va yangi imkoniyatlar ──

def test_pnl_excludes_non_halal_income_and_purification(app):
    F.record("kirim", 1_000_000, acc("Bank hisobi"), category_id=cat("Kurs to'lovlari").id)
    F.record("kirim", 300_000, acc("Bank hisobi"), category_id=cat("Boshqa kirim").id, shariat="nohalol")
    F.record("chiqim", 300_000, acc("Bank hisobi"), category_id=F.system_category("tozalash").id)
    db.session.commit()
    t = F.period_totals(date.today(), date.today())
    assert t == {"income": 1_000_000, "expense": 0, "profit": 1_000_000}
    assert F.pnl_table(date.today().year)["profit_total"] == 1_000_000


def test_edit_transaction_and_audit_log(app):
    from core.models import AuditLog
    client = app.test_client()
    login(client, "100100")
    tx = F.record("chiqim", 500_000, acc("Asosiy kassa"), category_id=cat("Ofis ijarasi").id)
    db.session.commit()
    client.post("/kassa/yangi", data={"tx_id": tx.id, "kind": "chiqim", "amount": "750 000",
                                      "account_id": acc("Bank hisobi").id, "category_id": cat("Ofis ijarasi").id,
                                      "date": date.today().isoformat(), "note": "tuzatildi"})
    tx = db.session.get(Transaction, tx.id)
    assert (tx.amount, tx.account.name, tx.note, tx.amount_uzs) == (750_000, "Bank hisobi", "tuzatildi", 750_000)
    assert Transaction.query.count() == 1
    assert AuditLog.query.filter_by(action="tx_update").count() == 1
    client.post(f"/kassa/{tx.id}/ochirish")
    assert AuditLog.query.filter_by(action="tx_delete").count() == 1
    assert "Kassa yozuvi tahrirlandi" in client.get("/sozlamalar/?bolim=tarix").get_data(as_text=True)


def test_manager_cannot_edit_or_change_programs(app):
    client = app.test_client()
    tx = F.record("kirim", 100, acc("Asosiy kassa"))
    db.session.commit()
    login(client, "300300")
    client.post("/kassa/yangi", data={"tx_id": tx.id, "kind": "kirim", "amount": "999", "account_id": tx.account_id})
    assert db.session.get(Transaction, tx.id).amount == 100
    assert client.post("/dasturlar/saqlash", data={"name": "X", "price": "1"}).status_code == 403
    assert client.get("/kassa/excel").status_code == 302


def test_deleting_repayment_reopens_debt(app):
    client = app.test_client()
    login(client, "100100")
    bank = acc("Bank hisobi")
    client.post("/qarz/yangi", data={"direction": "berilgan", "counterparty": "H", "amount": "1000", "account_id": bank.id})
    d = Debt.query.one()
    client.post(f"/qarz/{d.id}/qaytarish", data={"amount": "1000", "account_id": bank.id})
    assert db.session.get(Debt, d.id).closed
    repay = Transaction.query.filter_by(debt_id=d.id, kind="qarz_kirim").one()
    client.post(f"/kassa/{repay.id}/ochirish")
    assert not db.session.get(Debt, d.id).closed


def test_cross_site_post_rejected_and_headers(app):
    client = app.test_client()
    login(client, "100100")
    r = client.post("/kassa/yangi", data={"kind": "kirim", "amount": "1", "account_id": acc("Payme").id},
                    headers={"Origin": "https://evil.example"})
    assert r.status_code == 403 and Transaction.query.count() == 0
    r = client.post("/kassa/yangi", data={"kind": "kirim", "amount": "1", "account_id": acc("Payme").id},
                    headers={"Origin": "http://localhost"})
    assert Transaction.query.count() == 1
    h = client.get("/").headers
    assert h["X-Frame-Options"] == "DENY" and h["Cache-Control"] == "no-store"


def test_open_redirect_blocked():
    from views import is_local_path
    assert is_local_path("/kassa/?oy=2026-10")
    for bad in ("//evil.com", "/\\evil.com", "https://evil.com", "/\tx"):
        assert not is_local_path(bad)


def test_exports_are_valid_xlsx(app):
    from io import BytesIO
    from openpyxl import load_workbook
    from core.demo import demo
    demo()
    client = app.test_client()
    login(client, "100100")
    for url, sheets in (("/kassa/excel", 1), ("/sozlamalar/zaxira", 5), ("/hisobot/excel", 1)):
        r = client.get(url)
        assert r.status_code == 200, url
        wb = load_workbook(BytesIO(r.data))
        assert len(wb.sheetnames) == sheets and wb.worksheets[0].max_row > 1


def test_global_search_respects_roles(app):
    from core.demo import demo
    demo()
    client = app.test_client()
    login(client, "100100")
    page = client.get("/qidiruv/?q=Karim").get_data(as_text=True)
    assert "Karimov" in page and "Kassa yozuvlari" in page and "natija" in page
    login(client, "300300")
    page = client.get("/qidiruv/?q=Karim").get_data(as_text=True)
    assert "Qarzlar" not in page  # menejer qarz daftarini ko'rmaydi


# ── Kassa: Sheets uslubi ──

def test_ledger_running_balance_per_account_and_total(app):
    kassa, bank = acc("Asosiy kassa"), acc("Bank hisobi")
    month = date.today().replace(day=1)
    prev = month - timedelta(days=1)
    F.record("kirim", 1_000_000, kassa, date=prev)  # oldingi oy — boshlang'ich qoldiqqa kiradi
    a = F.record("kirim", 500_000, kassa, date=month)
    b = F.record("otkazma", 200_000, kassa, date=month, to_account_id=bank.id, to_amount=200_000)
    c = F.record("chiqim", 100_000, kassa, date=month)
    db.session.commit()
    rows = [a, b, c]
    opening, cur, flows = F.ledger_rows(rows, account=kassa, start=month)
    assert (opening, cur) == (1_000_000, "UZS")
    assert [flows[t.id]["bal"] for t in rows] == [1_500_000, 1_300_000, 1_200_000]
    # barcha hisoblar: o'tkazma umumiy qoldiqni o'zgartirmaydi va jamiga kirmaydi
    opening, cur, flows = F.ledger_rows(rows, start=month)
    assert [flows[t.id]["bal"] for t in rows] == [1_500_000, 1_500_000, 1_400_000]
    assert flows[b.id]["move"] and not flows[a.id]["move"]
    assert flows[c.id]["bal"] == F.total_balance_uzs()


def test_kassa_table_filters_and_hints(app):
    client = app.test_client()
    login(client, "100100")
    F.record("kirim", 750_000, acc("Asosiy kassa"), counterparty="Alisher Karimov", category_id=cat("Konsalting").id)
    F.record("chiqim", 90_000, acc("Asosiy kassa"), counterparty="Instagram", category_id=cat("Reklama (Instagram, Telegram)").id)
    db.session.commit()
    page = client.get("/kassa/").get_data(as_text=True)
    assert "Oy boshidagi qoldiq" in page and 'data-qe' in page
    assert '<option value="Alisher Karimov">' in page  # avto-to'ldirish ro'yxati
    page = client.get(f"/kassa/?kat={cat('Konsalting').id}").get_data(as_text=True)
    assert "Jami: 1 ta yozuv (filtr bo&#39;yicha)" in page or "Jami: 1 ta yozuv (filtr bo'yicha)" in page


def test_quick_entry_returns_focus_to_row(app):
    client = app.test_client()
    login(client, "100100")
    r = client.post("/kassa/yangi", data={"kind": "kirim", "amount": "1 200 000", "account_id": acc("Payme").id,
                                          "category_id": cat("Kurs to'lovlari").id, "counterparty": "Javohir",
                                          "date": date.today().isoformat(), "quick": "1", "next": "/kassa/?oy=2026-10"})
    tx = Transaction.query.one()
    assert tx.amount == 1_200_000 and tx.account.name == "Payme"
    assert r.headers["Location"].endswith(f"/kassa/?oy=2026-10&tez=1&belgi={tx.id}")


def test_sheet_parsers():
    from core.sheets import detect_header, parse_amount, parse_date, parse_kind
    assert parse_amount("1 500 000") == parse_amount("1,500,000") == parse_amount("1.500.000") == 1_500_000
    assert parse_amount("-250 000 so'm") == -250_000 and parse_amount("1500,50") == 1500.5 and parse_amount("abc") is None
    assert parse_date("05.10.2026") == parse_date("2026-10-05") == parse_date("5/10/26") == date(2026, 10, 5)
    assert parse_date("10/25/2026") == date(2026, 10, 25) and parse_date("31.02.2026") is None
    assert parse_kind("Kirim") == "kirim" and parse_kind("Xarajat") == "chiqim" and parse_kind("O'tkazma") == "skip"
    assert detect_header(["Sana", "Izoh", "Kirim", "Chiqim", "Qoldiq"]) == {0: "date", 1: "note", 2: "income", 3: "expense"}
    assert detect_header(["05.10.2026", "Kirim", "500000", "Asosiy kassa"]) is None


def test_import_from_sheets_preview_save_and_skip_duplicates(app):
    from core.models import AuditLog
    client = app.test_client()
    login(client, "100100")
    text = ("Sana\tKimdan / kimga\tKategoriya\tKirim\tChiqim\tQoldiq\tIzoh\n"
            "01.10.2026\tAlisher\tKurs to'lovlari\t1 500 000\t\t1 500 000\t3-potok\n"
            "02.10.2026\tInstagram\treklama\t\t450 000\t1 050 000\tOktabr\n"
            "03.10.2026\tNoma'lum\tYangi modda\t\t50 000\t1 000 000\t\n"
            "xato sana\tKimdir\t\t100\t\t\t\n")
    form = {"text": text, "default_account": acc("Asosiy kassa").id, "default_kind": "kirim"}
    page = client.post("/kassa/import", data=dict(form, action="preview")).get_data(as_text=True)
    assert "3 ta yangi" in page and "1 ta xato" in page and Transaction.query.count() == 0
    r = client.post("/kassa/import", data=dict(form, action="save"))
    assert r.status_code == 302 and "oy=2026-10" in r.headers["Location"]
    rows = Transaction.query.order_by(Transaction.date).all()
    assert [(t.kind, t.amount, t.counterparty) for t in rows] == [
        ("kirim", 1_500_000, "Alisher"), ("chiqim", 450_000, "Instagram"), ("chiqim", 50_000, "Noma'lum")]
    assert rows[1].category.name == "Reklama (Instagram, Telegram)"
    assert rows[2].category is None  # noma'lum kategoriya — "Boshqa" chiqim moddasi yo'q, kategoriyasiz qoladi
    assert AuditLog.query.filter_by(action="tx_import").count() == 1
    # qayta qo'yilsa — takror deb belgilanadi, ikki marta yozilmaydi
    page = client.post("/kassa/import", data=dict(form, action="save")).get_data(as_text=True)
    assert Transaction.query.count() == 3 and "3 ta bazada bor" in page


def test_import_xlsx_roundtrip_and_manager_blocked(app):
    from io import BytesIO
    client = app.test_client()
    login(client, "100100")
    tpl = client.get("/kassa/import/namuna.xlsx")
    assert tpl.status_code == 200
    r = client.post("/kassa/import", data={"action": "save", "default_account": acc("Asosiy kassa").id,
                                          "file": (BytesIO(tpl.data), "shablon.xlsx")},
                    content_type="multipart/form-data")
    assert r.status_code == 302 and Transaction.query.count() == 2
    login(client, "300300")
    assert client.get("/kassa/import").status_code == 302


# ── O'quvchi kartasi ──

def test_member_card_shows_full_history(app):
    from core.models import MemberEvent
    client = app.test_client()
    login(client, "100100")
    club = Program(name="Biznes klub", kind="klub", billing="oylik", price=1_000_000)
    course = Program(name="Kurs 9", kind="kurs", price=3_000_000)
    m = Member(full_name="Sardor Aliyev", phone="+998901112233", business="Savdo")
    db.session.add_all([club, course, m])
    db.session.commit()
    client.post(f"/dasturlar/{club.id}/yozish", data={"member_id": m.id, "joined_at": date.today().isoformat()})
    client.post(f"/azolar/{m.id}/yozish", data={"program_id": course.id, "price": "2 500 000"})
    e_club = Enrollment.query.filter_by(program_id=club.id).one()
    e_course = Enrollment.query.filter_by(program_id=course.id).one()
    assert e_course.price == 2_500_000 and e_course.member_id == m.id
    client.post("/kassa/yangi", data={"kind": "kirim", "amount": "1 000 000", "account_id": acc("Payme").id,
                                      "enrollment_id": e_club.id, "note": "1-oy"})
    client.post(f"/yozilish/{e_club.id}/uzaytirish")
    client.post(f"/yozilish/{e_course.id}/holat", data={"status": "tugatgan"})
    F.record("kirim", 200_000, acc("Asosiy kassa"), counterparty="sardor aliyev")  # dasturga bog'lanmagan
    db.session.add(Debt(direction="berilgan", counterparty="Sardor Aliyev", amount=500_000))
    db.session.commit()
    assert MemberEvent.query.filter_by(member_id=m.id).count() == 2
    page = client.get(f"/azolar/{m.id}").get_data(as_text=True)
    for text in ("Barcha tarix", "To&#39;lov qabul qilindi", "dasturiga yozildi", "uzaytirildi", "dasturni tugatdi",
                 "Qarz berildi", "200\u00a0000", "Dasturga yozish"):
        assert text in page or text.replace("&#39;", "'") in page, text
    # a'zolar paneli — kartochkalar, dastur sahifasida qator kartaga olib boradi
    page = client.get("/azolar").get_data(as_text=True)
    assert 'class="mcard' in page and f'/azolar/{m.id}' in page
    assert f'data-href="/azolar/{m.id}"' in client.get(f"/dasturlar/{club.id}").get_data(as_text=True)


def test_transfer_between_clubs_moves_payments(app):
    from core.models import AuditLog, MemberEvent
    client = app.test_client()
    login(client, "100100")
    a = Program(name="Klub A", kind="klub", billing="oylik", price=1_000_000)
    b = Program(name="Klub B", kind="klub", billing="oylik", price=1_500_000)
    m = Member(full_name="Jasur Karimov")
    db.session.add_all([a, b, m])
    db.session.commit()
    client.post(f"/dasturlar/{a.id}/yozish", data={"member_id": m.id})
    e = Enrollment.query.filter_by(program_id=a.id).one()
    client.post("/kassa/yangi", data={"kind": "kirim", "amount": "1 000 000", "account_id": acc("Payme").id,
                                      "enrollment_id": e.id})
    total_before = F.total_balance_uzs()
    r = client.post("/yozilish/otkazish", data={"eid": e.id, "program_id": b.id, "price": "1 500 000",
                                                "date": date.today().isoformat(), "note": "vaqti mos emas"})
    assert r.status_code == 302 and r.headers["Location"].endswith(f"/azolar/{m.id}")
    e = db.session.get(Enrollment, e.id)
    new = Enrollment.query.filter_by(program_id=b.id).one()
    assert e.status == "chiqqan" and new.status == "faol" and new.price == 1_500_000 and new.valid_until
    paid_map = F.enrollment_paid_map()
    assert F.enrollment_row(new, paid_map)["due"] == 500_000          # 1.5 mln − o'tgan 1 mln
    assert F.program_stats(a)["collected"] == 0 and F.program_stats(b)["collected"] == 1_000_000
    assert F.total_balance_uzs() == total_before                      # kassa o'zgarmaydi
    assert MemberEvent.query.filter_by(kind="transfer").count() == 1
    assert AuditLog.query.filter_by(action="transfer").count() == 1
    page = client.get(f"/azolar/{m.id}").get_data(as_text=True)
    assert "ga o&#39;tkazildi" in page or "ga o'tkazildi" in page
    # yopilgan yozilishni qayta o'tkazib bo'lmaydi; menejerga ruxsat yo'q
    client.post("/yozilish/otkazish", data={"eid": e.id, "program_id": b.id})
    assert Enrollment.query.count() == 2
    login(client, "300300")
    client.post("/yozilish/otkazish", data={"eid": new.id, "program_id": a.id})
    assert db.session.get(Enrollment, new.id).status == "faol"


def test_refund_reduces_income_not_expense(app):
    from core.models import AuditLog
    client = app.test_client()
    login(client, "100100")
    p = Program(name="Kurs R", kind="kurs", price=4_000_000)
    m = Member(full_name="Bobur Aliyev")
    db.session.add_all([p, m])
    db.session.commit()
    client.post(f"/dasturlar/{p.id}/yozish", data={"member_id": m.id})
    e = Enrollment.query.one()
    payme = acc("Payme")
    client.post("/kassa/yangi", data={"kind": "kirim", "amount": "4 000 000", "account_id": payme.id,
                                      "category_id": cat("Kurs to'lovlari").id, "enrollment_id": e.id})
    bal0, today = F.account_balances()[payme.id], date.today()
    # to'langanidan ko'p qaytarib bo'lmaydi
    client.post("/yozilish/qaytarish", data={"eid": e.id, "amount": "5 000 000", "account_id": payme.id})
    assert Transaction.query.count() == 1
    # qisman qaytarish — chegirma: narx kamayadi, qarz paydo bo'lmaydi
    client.post("/yozilish/qaytarish", data={"eid": e.id, "amount": "1 000 000", "account_id": payme.id,
                                             "then": "chegirma", "note": "chegirma berildi"})
    e = db.session.get(Enrollment, e.id)
    row = F.enrollment_row(e)
    assert (e.price, row["paid"], row["due"], e.status) == (3_000_000, 3_000_000, 0, "faol")
    t = F.period_totals(today, today)
    assert t["income"] == 3_000_000 and t["expense"] == 0           # xarajat emas — kirim kamaydi
    assert F.account_balances()[payme.id] == bal0 - 1_000_000
    assert F.program_stats(p)["collected"] == 3_000_000
    # kassa daftarida chiqim ustunida
    refund_tx = Transaction.query.filter(Transaction.amount < 0).one()
    _, _, flows = F.ledger_rows([refund_tx], account=payme, start=today.replace(day=1))
    assert (flows[refund_tx.id]["inc"], flows[refund_tx.id]["out"]) == (0, 1_000_000)
    # to'liq qaytarish — dasturdan chiqadi
    client.post("/yozilish/qaytarish", data={"eid": e.id, "amount": "3 000 000", "account_id": payme.id, "then": "chiqdi"})
    assert db.session.get(Enrollment, e.id).status == "chiqqan" and F.program_stats(p)["collected"] == 0
    assert AuditLog.query.filter_by(action="refund").count() == 2
    page = client.get(f"/azolar/{m.id}").get_data(as_text=True)
    assert "Pul qaytarildi" in page and "Qaytarildi" in client.get("/kassa/").get_data(as_text=True)
    login(client, "300300")
    client.post("/yozilish/qaytarish", data={"eid": e.id, "amount": "1", "account_id": payme.id})
    assert Transaction.query.filter(Transaction.amount < 0).count() == 2


# ── Kassa auditi ──

def test_kassa_kirim_linked_to_student_reduces_debt(app):
    client = app.test_client()
    login(client, "100100")
    p = Program(name="Kurs L", kind="kurs", price=2_000_000)
    m = Member(full_name="Otabek Nazarov")
    db.session.add_all([p, m])
    db.session.commit()
    client.post(f"/dasturlar/{p.id}/yozish", data={"member_id": m.id})
    e = Enrollment.query.one()
    page = client.get("/kassa/").get_data(as_text=True)
    assert "data-members=" in page and '<option value="Otabek Nazarov">' in page  # ism ro'yxatda, to'lovsiz ham
    # kassadan (tez yozish qatori) o'quvchi to'loviga bog'lab kirim
    client.post("/kassa/yangi", data={"kind": "kirim", "amount": "800 000", "account_id": acc("Click").id,
                                      "counterparty": "Otabek Nazarov", "enrollment_id": e.id, "quick": "1"})
    assert F.enrollment_row(db.session.get(Enrollment, e.id))["due"] == 1_200_000
    tx = Transaction.query.one()
    assert tx.program_id == p.id


def test_kassa_warnings_duplicate_negative_future(app):
    client = app.test_client()
    login(client, "100100")
    data = {"kind": "chiqim", "amount": "300 000", "account_id": acc("Asosiy kassa").id, "counterparty": "Zal",
            "date": date.today().isoformat()}
    client.post("/kassa/yangi", data=data)
    client.post("/kassa/yangi", data=data)
    with client.session_transaction() as sess:
        msgs = [m for cat_, m in sess.get("_flashes", []) if cat_ == "warn"]
    assert any("xuddi shunday yozuv" in m for m in msgs) and any("manfiy" in m for m in msgs)
    client.post("/kassa/yangi", data=dict(data, kind="kirim", amount="1000",
                                          date=(date.today() + timedelta(days=3)).isoformat()))
    with client.session_transaction() as sess:
        assert any("kelajakda" in m for c, m in sess.get("_flashes", []) if c == "warn")


def test_kassa_bulk_category_delete_and_uncategorized_filter(app):
    client = app.test_client()
    login(client, "100100")
    a = F.record("chiqim", 100, acc("Asosiy kassa"))
    b = F.record("chiqim", 200, acc("Asosiy kassa"))
    c = F.record("kirim", 300, acc("Asosiy kassa"))
    db.session.commit()
    page = client.get("/kassa/?kat=yoq").get_data(as_text=True)
    assert "Kategoriyasiz (3)" in page and "Jami: 3 ta yozuv" in page
    office = cat("Ofis ijarasi")
    client.post("/kassa/ommaviy", data={"ids": [a.id, b.id, c.id], "action": "category", "category_id": office.id})
    assert {t.id for t in Transaction.query.filter_by(category_id=office.id)} == {a.id, b.id}  # kirim o'zgarmaydi
    client.post("/kassa/ommaviy", data={"ids": [a.id, c.id], "action": "delete"})
    assert [t.id for t in Transaction.query.all()] == [b.id]
    login(client, "300300")
    client.post("/kassa/ommaviy", data={"ids": [b.id], "action": "delete"})
    assert Transaction.query.count() == 1


def test_kassa_export_has_balance_and_reimports_as_duplicates(app):
    from io import BytesIO
    from openpyxl import load_workbook
    client = app.test_client()
    login(client, "100100")
    kassa = acc("Asosiy kassa")
    F.record("kirim", 1_000_000, kassa, counterparty="Ali", category_id=cat("Konsalting").id)
    F.record("chiqim", 250_000, kassa, counterparty="Zal", category_id=cat("Zal ijarasi").id)
    db.session.commit()
    r = client.get(f"/kassa/excel?hisob={kassa.id}")
    ws = load_workbook(BytesIO(r.data)).active
    head = [c.value for c in ws[1]]
    assert head[7:10] == ["Kirim, so'm", "Chiqim, so'm", "Qoldiq, so'm"]
    assert ws.cell(4, 10).value == 750_000 and str(ws.cell(5, 8).value).startswith("=SUM(")
    r = client.post("/kassa/import", data={"action": "preview", "file": (BytesIO(r.data), "kassa.xlsx")},
                    content_type="multipart/form-data")
    assert "2 ta bazada bor" in r.get_data(as_text=True)


# ── Dars jadvali (kalendar) ──

def test_calendar_helpers():
    from datetime import time
    from core import calendar as C
    from core.models import Lesson
    start = date(2026, 10, 5)  # dushanba
    days = C.expand_weekly(start, [1, 3], count=5)
    assert [d.weekday() for d in days] == [1, 3, 1, 3, 1] and days[0] == date(2026, 10, 6)
    assert C.expand_weekly(start, [0], until=date(2026, 10, 26)) == [date(2026, 10, d) for d in (5, 12, 19, 26)]
    a = Lesson(start=time(9, 0), end=time(11, 0))
    b = Lesson(start=time(10, 0), end=time(12, 0))
    c = Lesson(start=time(13, 0), end=time(14, 0))
    lay = {id(x["lesson"]): x for x in C.layout_day([a, b, c], 8 * 60)}
    assert (lay[id(a)]["col"], lay[id(b)]["col"], lay[id(a)]["cols"]) == (0, 1, 2)  # kesishganlar yonma-yon
    assert (lay[id(c)]["cols"], lay[id(c)]["top"]) == (1, 300)
    assert C.fmt_range(date(2026, 9, 28), date(2026, 10, 4)) == "28 sentabr – 4 oktabr 2026"


def test_calendar_series_edit_move_delete(app):
    from core.models import Lesson
    client = app.test_client()
    login(client, "300300")  # menejer ham jadval tuza oladi
    p = Program(name="Kurs K", kind="kurs", price=1)
    db.session.add(p)
    db.session.commit()
    r = client.post("/kalendar/saqlash", data={"program_id": p.id, "date": "2026-10-06", "start": "19:00", "end": "21:00",
                                               "repeat": "weekly", "weekdays": ["1", "3"], "count": "6",
                                               "location": "Asosiy zal", "teacher": "Ustoz A"})
    assert r.status_code == 302
    les = Lesson.query.order_by(Lesson.date).all()
    assert len(les) == 6 and les[0].title == "1-dars" and les[5].title == "6-dars"
    assert len({x.series_id for x in les}) == 1 and les[1].date == date(2026, 10, 8)
    # xato vaqt
    client.post("/kalendar/saqlash", data={"program_id": p.id, "date": "2026-10-06", "start": "20:00", "end": "19:00"})
    assert Lesson.query.count() == 6
    # seriyani 3-darsdan boshlab tahrirlash — oldingilari o'zgarmaydi
    third = les[2]
    client.post("/kalendar/saqlash", data={"lesson_id": third.id, "scope": "series", "program_id": p.id,
                                           "date": third.date.isoformat(), "start": "18:00", "end": "20:00",
                                           "location": "Kichik zal", "title": third.title})
    les = Lesson.query.order_by(Lesson.date).all()
    assert [x.start.hour for x in les] == [19, 19, 18, 18, 18, 18] and les[5].location == "Kichik zal"
    # to'qnashuv: shu zalda shu vaqtga boshqa dars — ogohlantirish
    client.post("/kalendar/saqlash", data={"program_id": p.id, "date": les[3].date.isoformat(), "start": "19:00",
                                           "end": "20:30", "location": "kichik zal", "title": "Qo'shimcha"})
    with client.session_transaction() as sess:
        assert any("To'qnashuv" in m for c, m in sess.get("_flashes", []) if c == "warn")
    # sudrab ko'chirish (JSON)
    r = client.post(f"/kalendar/{les[0].id}/kochirish", json={"date": "2026-10-07", "start": "10:00", "end": "11:30"})
    assert r.get_json()["ok"] and db.session.get(Lesson, les[0].id).date == date(2026, 10, 7)
    assert client.post(f"/kalendar/{les[0].id}/kochirish", json={"date": "x"}).status_code == 400
    client.post(f"/kalendar/{les[1].id}/holat", data={"status": "otdi"})
    assert db.session.get(Lesson, les[1].id).status == "otdi"
    # seriyani 5-darsdan boshlab o'chirish
    client.post(f"/kalendar/{les[4].id}/ochirish", data={"scope": "series"})
    assert Lesson.query.filter(Lesson.series_id.isnot(None)).count() == 4
    for v in ("hafta", "soat", "kun", "oy", "yol", "royxat"):
        page = client.get(f"/kalendar/?v={v}&d=2026-10-08").get_data(as_text=True)
        assert "Yangi dars" in page, v
    assert "Dars jadvali" in client.get(f"/dasturlar/{p.id}").get_data(as_text=True)


# ── Davomat va QR skaner ──

def test_member_code_roundtrip_and_typos():
    from core.attendance import member_code, parse_code
    for mid in (1, 23, 4567, 999999):
        code = member_code(mid)
        assert len(code) == 9 and parse_code(code) == mid and parse_code(f"  {code.lower()}\n") == mid
        bad = code[:-1] + str((int(code[-1]) + 1) % 10)
        assert parse_code(bad) is None  # nazorat raqami xatoni ushlaydi
    assert parse_code("https://assoodiq.uz/a/" + member_code(5)) == 5 and parse_code("salom") is None


def test_attendance_scan_mark_and_stats(app):
    from datetime import time
    from core.attendance import member_code
    from core.models import Attendance, Lesson
    client = app.test_client()
    login(client, "300300")  # menejer ham davomat oladi
    p = Program(name="Kurs D", kind="kurs", price=1)
    a, b, c = Member(full_name="Ali Valiyev"), Member(full_name="Bobur Sodiqov"), Member(full_name="Mehmon Bek")
    db.session.add_all([p, a, b, c])
    db.session.commit()
    for m in (a, b):
        db.session.add(Enrollment(member_id=m.id, program_id=p.id, price=1, joined_at=date(2026, 1, 1)))
    les = Lesson(program_id=p.id, title="1-dars", date=date.today(), start=time(0, 0), end=time(23, 59))
    db.session.add(les)
    db.session.commit()
    page = client.get(f"/kalendar/{les.id}/davomat").get_data(as_text=True)
    assert "Ali Valiyev" in page and "Bobur Sodiqov" in page and "Skaner" in page
    # skaner: kod → keldi/kechikdi; ikkinchi marta — takror
    r = client.post(f"/kalendar/{les.id}/davomat/skaner", json={"code": member_code(a.id)}).get_json()
    assert r["ok"] and r["status"] in ("keldi", "kechikdi") and not r["already"] and r["stats"]["present"] == 1
    assert client.post(f"/kalendar/{les.id}/davomat/skaner", json={"code": member_code(a.id)}).get_json()["already"]
    assert db.session.get(Lesson, les.id).status == "otdi"
    # noto'g'ri kod va dasturda yo'q a'zo
    assert not client.post(f"/kalendar/{les.id}/davomat/skaner", json={"code": "AQ0000010"}).get_json()["ok"]
    r = client.post(f"/kalendar/{les.id}/davomat/skaner", json={"code": member_code(c.id)}).get_json()
    assert not r["ok"] and r["need_force"]
    r = client.post(f"/kalendar/{les.id}/davomat/skaner", json={"code": member_code(c.id), "force": True}).get_json()
    assert r["ok"] and r["stats"]["total"] == 3  # mehmon ham ro'yxatga qo'shildi
    # qo'lda: kelmadi, keyin belgini olib tashlash
    r = client.post(f"/kalendar/{les.id}/davomat/belgi", json={"member_id": b.id, "status": "kelmadi"}).get_json()
    assert r["stats"]["kelmadi"] == 1
    r = client.post(f"/kalendar/{les.id}/davomat/belgi", json={"member_id": b.id, "status": ""}).get_json()
    assert r["stats"]["kelmadi"] == 0 and r["stats"]["unmarked"] == 1
    client.post(f"/kalendar/{les.id}/davomat/hammasi", data={"status": "keldi"})
    present = Attendance.query.filter(Attendance.lesson_id == les.id, Attendance.status.in_(("keldi", "kechikdi")))
    assert present.count() == 3  # Ali (skaner), mehmon (skaner) va Bobur («qolganlar keldi»)
    # kartalar va o'quvchi kartasidagi davomat
    page = client.get(f"/azolar/kartalar?dastur={p.id}").get_data(as_text=True)
    assert member_code(a.id) in page and "<svg" in page and "Bobur Sodiqov" in page
    assert "Davomat:" in client.get(f"/azolar/{a.id}").get_data(as_text=True)
    # dars o'chirilsa davomati ham o'chadi
    client.post(f"/kalendar/{les.id}/ochirish")
    assert Attendance.query.count() == 0


def test_dashboard_today_lessons_states(app):
    from datetime import datetime, time
    from core.attendance import mark, today_lessons
    from core.models import Lesson
    p = Program(name="Kurs B", kind="kurs", price=1)
    m = Member(full_name="Sanjar")
    db.session.add_all([p, m])
    db.session.commit()
    db.session.add(Enrollment(member_id=m.id, program_id=p.id, price=1, joined_at=date(2026, 1, 1)))
    d = date.today()
    mk = lambda h0, h1, **kw: Lesson(program_id=p.id, title=f"{h0}-dars", date=d, start=time(h0), end=time(h1), **kw)
    done, nomark, live, soon, cancel = mk(8, 9), mk(9, 10), mk(11, 13), mk(15, 17), mk(18, 19, status="bekor")
    db.session.add_all([done, nomark, live, soon, cancel])
    db.session.commit()
    mark(done, m.id, "keldi")
    db.session.commit()
    lt = today_lessons(now=datetime.combine(d, time(12, 0)))
    states = {x["lesson"].title: (x["state"], x["label"]) for x in lt["lessons"]}
    assert states["8-dars"] == ("done", "1 / 1 keldi") and states["9-dars"][0] == "nomark"
    assert states["11-dars"][0] == "live" and states["15-dars"] == ("soon", "3 soat 0 daq dan keyin")
    assert states["18-dars"][0] == "bekor" and (lt["total"], lt["done"], lt["nomark"]) == (4, 2, 1)
    client = app.test_client()
    login(client, "100100")
    page = client.get("/").get_data(as_text=True)
    assert "Bugungi darslar" in page and "Davomat" in page
    # dars bo'lmagan kuni blok chiqmaydi
    from core.models import Attendance
    Attendance.query.delete()
    Lesson.query.delete()
    db.session.commit()
    assert "Bugungi darslar" not in client.get("/").get_data(as_text=True)


# ── Telegram ──

@pytest.fixture
def tg(app, monkeypatch):
    """Telegram API o'rniga yozib boruvchi soxta funksiya."""
    from core import telegram as T
    calls = []

    def fake_api(method, payload=None, tok=None, timeout=10):
        calls.append((method, payload or {}))
        if method == "getMe":
            return True, {"username": "assoodiq_test_bot"}
        if method == "sendMessage" and str(payload.get("chat_id")) == "999":
            return False, "Forbidden: bot was blocked by the user"
        return True, {"message_id": len(calls)}

    monkeypatch.setattr(T, "api", fake_api)
    F.set_setting("tg_token", "123:ABC")
    F.set_setting("tg_bot", "assoodiq_test_bot")
    db.session.commit()
    return calls


_UPD_IDS = __import__("itertools").count(1000)


def _upd(chat_id, text=None, **msg):
    m = {"chat": {"id": chat_id, "type": "private"}, "from": {"id": chat_id, "username": "ali"}, **msg}
    if text is not None:
        m["text"] = text
    return {"update_id": next(_UPD_IDS), "message": m}


def _sent(calls):
    return [p for meth, p in calls if meth == "sendMessage"]


def test_telegram_webhook_secret_and_linking(app, tg):
    from core import telegram as T
    from core.attendance import member_code
    from core.models import TelegramLink
    client = app.test_client()
    a = Member(full_name="Ali Valiyev", phone="+998 90 123 45 67")
    b = Member(full_name="Bobur Sodiqov", phone="90 765 43 21")
    db.session.add_all([a, b])
    db.session.commit()
    secret = T.webhook_secret()
    hdr = {"X-Telegram-Bot-Api-Secret-Token": secret}
    # kalitsiz yoki noto'g'ri kalit — 404
    assert client.post(f"/telegram/hook/{secret}", json=_upd(1, "/start")).status_code == 404
    assert client.post("/telegram/hook/wrong", json=_upd(1, "/start"), headers=hdr).status_code == 404
    # /start AQ… — shaxsiy havola orqali ulanish
    r = client.post(f"/telegram/hook/{secret}", json=_upd(111, f"/start {member_code(a.id)}"), headers=hdr)
    assert r.status_code == 200
    link = TelegramLink.query.filter_by(member_id=a.id).one()
    assert link.chat_id == "111" and link.active and link.username == "ali"
    assert "Ali Valiyev" in _sent(tg)[-1]["text"]
    # telefon raqami orqali ulanish (oxirgi 9 raqam bo'yicha)
    client.post(f"/telegram/hook/{secret}", json=_upd(222, contact={"phone_number": "998907654321", "user_id": 222}), headers=hdr)
    assert TelegramLink.query.filter_by(member_id=b.id).one().chat_id == "222"
    client.post(f"/telegram/hook/{secret}", json=_upd(333, contact={"phone_number": "998900000000", "user_id": 333}), headers=hdr)
    assert "topilmadi" in _sent(tg)[-1]["text"]
    # bot tugmalari: jadval va to'lov
    p = Program(name="Kurs T", kind="kurs", price=500_000)
    db.session.add(p)
    db.session.commit()
    db.session.add(Enrollment(member_id=a.id, program_id=p.id, price=500_000, joined_at=date(2026, 1, 1)))
    db.session.commit()
    client.post(f"/telegram/hook/{secret}", json=_upd(111, "💳 To'lov"), headers=hdr)
    assert "500 000" in _sent(tg)[-1]["text"] or "500 000" in _sent(tg)[-1]["text"]
    client.post(f"/telegram/hook/{secret}", json=_upd(111, "📅 Jadval"), headers=hdr)
    assert "dars" in _sent(tg)[-1]["text"].lower()
    # /stop — eslatmalar to'xtaydi
    client.post(f"/telegram/hook/{secret}", json=_upd(111, "/stop"), headers=hdr)
    assert not TelegramLink.query.filter_by(member_id=a.id).one().active


def test_telegram_auto_reminders_sent_once(app, tg):
    from datetime import datetime, time
    from core import telegram as T
    from core.models import Lesson, Notification, TelegramLink
    p = Program(name="Kurs R", kind="kurs", price=300_000)
    a, b, c = Member(full_name="Ali Valiyev"), Member(full_name="Bobur"), Member(full_name="Blok Bek")
    db.session.add_all([p, a, b, c])
    db.session.commit()
    for m in (a, b, c):
        db.session.add(Enrollment(member_id=m.id, program_id=p.id, price=300_000, joined_at=date(2026, 1, 1)))
    db.session.add_all([TelegramLink(member_id=a.id, chat_id="111"), TelegramLink(member_id=c.id, chat_id="999")])
    d = date.today() + timedelta(days=1)
    les = Lesson(program_id=p.id, title="3-dars", date=d, start=time(10, 0), end=time(12, 0), location="Zal A")
    db.session.add(les)
    db.session.commit()
    now = datetime.combine(d, time(8, 0))  # 2 soat oldin (standart: 3 soat)
    T.run_due(now)
    texts = [x["text"] for x in _sent(tg) if x["chat_id"] == "111"]
    assert len(texts) == 1 and "3-dars" in texts[0] and "10:00" in texts[0]
    # Blok Bek botni bloklagan — havola o'chadi, xato yoziladi
    assert not TelegramLink.query.filter_by(member_id=c.id).one().active
    assert Notification.query.filter_by(status="failed").count() == 1
    # qayta ishga tushsa — takror yuborilmaydi
    n = len(_sent(tg))
    T.run_due(now)
    assert len(_sent(tg)) == n
    # oylik to'lov eslatmasi (yoqilgan bo'lsa, belgilangan kuni)
    F.set_setting("tg_auto_pay", "1")
    F.set_setting("tg_pay_day", str(d.day))
    db.session.commit()
    T.run_due(datetime.combine(d, time(11, 0)))
    pay = [x["text"] for x in _sent(tg)[n:]]
    assert len(pay) == 1 and "300" in pay[0]
    T.run_due(datetime.combine(d, time(11, 5)))
    assert len(_sent(tg)) == n + 1


def test_telegram_settings_broadcast_and_lesson_reminder(app, tg):
    from datetime import time
    from core.models import Lesson, Notification, TelegramLink
    client = app.test_client()
    login(client, "100100")
    page = client.get("/sozlamalar/?bolim=telegram").get_data(as_text=True)
    assert "@assoodiq_test_bot" in page and "Botni saytga ulash" in page
    # yangi token getMe orqali tekshiriladi
    client.post("/telegram/sozla", data={"token": "456:XYZ", "tg_auto_lessons": "on", "tg_lead_hours": "2",
                                         "tg_expiry_days": "5", "tg_pay_day": "40"})
    assert F.get_setting("tg_token") == "456:XYZ" and F.get_setting("tg_lead_hours") == "2"
    assert F.get_setting("tg_pay_day") == "28" and F.get_setting("tg_auto_expiry") == "0"
    assert ("getMe", {}) in tg
    client.post("/telegram/webhook", data={"base": "http://assoodiq.up.railway.app"})
    hook = next(p for meth, p in tg if meth == "setWebhook")
    assert hook["url"].startswith("https://assoodiq.up.railway.app/telegram/hook/") and hook["secret_token"]
    # dastur a'zolariga xabar: ulanganlarga boradi, qolganlar sanaladi
    p = Program(name="Klub X", kind="klub", price=100_000)
    a, b = Member(full_name="Ali Valiyev"), Member(full_name="Bobur")
    db.session.add_all([p, a, b])
    db.session.commit()
    for m in (a, b):
        db.session.add(Enrollment(member_id=m.id, program_id=p.id, price=100_000, joined_at=date(2026, 1, 1)))
    db.session.add(TelegramLink(member_id=a.id, chat_id="111"))
    db.session.commit()
    r = client.post("/telegram/yuborish", data={"scope": "dastur", "program_id": p.id,
                                                "text": "Salom, {ism}! <script> {dastur}"}, follow_redirects=True)
    assert r.status_code == 200
    msg = _sent(tg)[-1]
    assert msg["chat_id"] == "111" and msg["text"] == "Salom, Ali! &lt;script&gt; Klub X"
    assert Notification.query.filter_by(member_id=b.id).count() == 0
    # qarzdorlarga — shablon bilan
    client.post("/telegram/yuborish", data={"scope": "qarzdorlar", "program_id": p.id})
    assert "100" in _sent(tg)[-1]["text"]
    # darsdan eslatma tugmasi va o'quvchi kartasi
    les = Lesson(program_id=p.id, title="Uchrashuv", date=date.today() + timedelta(days=2), start=time(18), end=time(19))
    db.session.add(les)
    db.session.commit()
    client.post(f"/telegram/dars/{les.id}")
    assert "Uchrashuv" in _sent(tg)[-1]["text"]
    client.post(f"/telegram/dars/{les.id}")  # takror bosilsa — ikkinchi marta yuborilmaydi
    assert sum("Uchrashuv" in x["text"] for x in _sent(tg)) == 1
    page = client.get(f"/azolar/{b.id}").get_data(as_text=True)
    assert "t.me/assoodiq_test_bot?start=" in page
    assert "Ulangan" in client.get(f"/azolar/{a.id}").get_data(as_text=True)


# ── O'quvchi ilovasi ──

def test_prayer_times_and_qibla():
    from datetime import datetime
    from core import prayer as P
    lat, lng = P.CITIES["Toshkent"]
    for d in (date(2026, 3, 20), date(2026, 6, 21), date(2026, 12, 21)):
        t = P.compute(d, lat, lng)
        order = [t[k] for k, _ in P.PRAYERS]
        assert all(order[i] < order[i + 1] for i in range(5)), (d, order)
        assert 12 <= t["peshin"] <= 12.75  # Toshkentda quyosh tikkasi ~12:10–12:35
        assert P.compute(d, lat, lng, asr="shafii")["asr"] < t["asr"]  # Hanafiy asr kechroq
    # tuzatish daqiqalari
    base = P.compute(date(2026, 3, 20), lat, lng)["bomdod"]
    assert abs(P.compute(date(2026, 3, 20), lat, lng, offsets=[5, 0, 0, 0, 0, 0])["bomdod"] - base - 5 / 60) < 1e-9
    deg, km = P.qibla(lat, lng)
    assert 238 <= deg <= 243 and 3400 <= km <= 3650 and P.compass_word(deg) == "janubi-g'arb"
    key, name, at = P.next_prayer(datetime(2026, 10, 11, 23, 30), lat, lng)
    assert key == "bomdod" and at.date() == date(2026, 10, 12)
    assert P.iso_at(datetime(2026, 1, 2, 4, 5)).endswith("+05:00")


def _course(with_valid_until=None):
    from core.models import Lecture, Module
    p = Program(name="Kurs A", kind="kurs", price=1_000_000)
    m = Member(full_name="Ali Valiyev", phone="+998 90 111 22 33")
    db.session.add_all([p, m])
    db.session.commit()
    e = Enrollment(member_id=m.id, program_id=p.id, price=1_000_000, joined_at=date(2026, 1, 1),
                   valid_until=with_valid_until)
    mod = Module(program_id=p.id, title="1-modul", position=0)
    db.session.add_all([e, mod])
    db.session.commit()
    l1 = Lecture(module_id=mod.id, title="Kirish", kind="video", duration=600, position=0, free=True)
    l2 = Lecture(module_id=mod.id, title="Asosiy dars", kind="video", duration=1200, position=1)
    l3 = Lecture(module_id=mod.id, title="Qo'llanma", kind="matn", body="Birinchi xatboshi.\n\nIkkinchi.", position=2)
    db.session.add_all([l1, l2, l3])
    db.session.commit()
    return p, m, e, (l1, l2, l3)


def test_student_invite_login_courses_and_progress(app):
    from core import student as S
    from core.models import LectureProgress
    p, m, e, (l1, l2, l3) = _course()
    client = app.test_client()
    # kirmagan — kirish sahifasiga; xodim sahifalari ham yopiq
    assert client.get("/app/").headers["Location"].endswith("/app/kirish")
    tok = S.issue_invite(m)
    assert client.get(f"/app/k/{tok}").headers["Location"].endswith("/app/")
    home = client.get("/app/").get_data(as_text=True)
    assert "Ali" in home and "Kurs A" in home and "Keyingi namoz" in home
    assert client.get("/").status_code == 302  # o'quvchi sessiyasi xodim paneliga kiritmaydi
    for path in ("/app/kurslar", f"/app/kurs/{p.id}", "/app/jadval", "/app/ibodat", "/app/zakot", "/app/profil"):
        assert client.get(path).status_code == 200, path
    page = client.get(f"/app/dars/{l3.id}").get_data(as_text=True)
    assert "Birinchi xatboshi." in page and "Ikkinchi." in page
    # 90% ko'rilsa — tugatilgan
    r = client.post(f"/app/dars/{l2.id}/progress", json={"position": 300, "duration": 1200}).get_json()
    assert r["ok"] and not r["done"]
    assert client.post(f"/app/dars/{l2.id}/progress", json={"position": 1100, "duration": 1200}).get_json()["done"]
    client.post(f"/app/dars/{l3.id}/progress", json={"done": True})
    c = S.courses(m)[0]
    assert (c["done"], c["lectures"], c["next"].id) == (2, 3, l1.id)
    assert "AQ" in client.get("/app/profil").get_data(as_text=True)  # QR karta kodi
    # yangi havola — eskisi ishlamaydi; barcha qurilmalardan chiqarish
    tok2 = S.issue_invite(m)
    assert client.get(f"/app/k/{tok}").headers["Location"].endswith("/app/kirish")
    S.revoke_all(m)
    assert client.get("/app/").headers["Location"].endswith("/app/kirish")
    assert client.get(f"/app/k/{tok2}").headers["Location"].endswith("/app/kirish")
    assert LectureProgress.query.filter_by(member_id=m.id, done=True).count() == 2


def test_student_locked_when_subscription_expired(app):
    from core import student as S
    p, m, e, (l1, l2, l3) = _course(with_valid_until=date.today() - timedelta(days=1))
    other = Program(name="Boshqa kurs", kind="kurs", price=1)
    db.session.add(other)
    db.session.commit()
    client = app.test_client()
    client.get(f"/app/k/{S.issue_invite(m)}")
    assert S.can_view(m, l1) == (True, "")          # ochiq dars
    assert S.can_view(m, l2) == (False, "expired")
    assert "Dars yopiq" in client.get(f"/app/dars/{l2.id}").get_data(as_text=True)
    assert client.post(f"/app/dars/{l2.id}/progress", json={"done": True}).status_code == 403
    assert "Obunangiz" in client.get(f"/app/kurs/{p.id}").get_data(as_text=True)
    assert client.get(f"/app/kurs/{other.id}").status_code == 302  # yozilmagan kurs


def test_student_login_by_telegram_code_and_mini_app(app, tg):
    import hashlib
    import hmac
    import json as _json
    import time as _time
    from urllib.parse import urlencode
    from core.models import TelegramLink
    p, m, e, _ = _course()
    db.session.add(TelegramLink(member_id=m.id, chat_id="777"))
    db.session.commit()
    client = app.test_client()
    # noma'lum raqam
    page = client.post("/app/kirish", data={"phone": "+998 99 000 00 00"}).get_data(as_text=True)
    assert "Kod yuborib bo" in page
    page = client.post("/app/kirish", data={"phone": "90 111 22 33"}).get_data(as_text=True)
    assert "Kodni kiriting" in page
    msg = _sent(tg)[-1]
    assert msg["chat_id"] == "777"
    import re
    code = re.search(r"<b>(\d{6})</b>", msg["text"]).group(1)
    assert "Kod noto" in client.post("/app/kirish/kod", data={"code": "000000" if code != "000000" else "111111"}).get_data(as_text=True)
    assert client.post("/app/kirish/kod", data={"code": code}).headers["Location"].endswith("/app/")
    assert client.get("/app/").status_code == 200
    # Telegram Mini App: imzolangan initData bilan avtomatik kirish
    token = F.get_setting("tg_token")
    fields = {"auth_date": str(int(_time.time())), "query_id": "AAE", "user": _json.dumps({"id": 777, "first_name": "Ali"})}
    check = "\n".join(f"{k}={v}" for k, v in sorted(fields.items()))
    secret = hmac.new(b"WebAppData", token.encode(), hashlib.sha256).digest()
    fields["hash"] = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
    c2 = app.test_client()
    r = c2.post("/app/tg", json={"initData": urlencode(fields)}).get_json()
    assert r["ok"] and c2.get("/app/").status_code == 200
    bad = dict(fields, user=_json.dumps({"id": 778, "first_name": "X"}))
    assert app.test_client().post("/app/tg", json={"initData": urlencode(bad)}).status_code == 400


def test_curriculum_editor_bunny_upload_and_settings(app, monkeypatch):
    from core import bunny
    from core.models import Lecture, Module
    calls = []

    def fake_api(method, path, payload=None, timeout=15):
        calls.append((method, path, payload))
        if method == "POST" and path == "/videos":
            return True, {"guid": "vid-123"}
        if method == "GET" and path.startswith("/videos/"):
            return True, {"status": 4, "length": 754, "encodeProgress": 100}
        return True, {}

    monkeypatch.setattr(bunny, "api", fake_api)
    client = app.test_client()
    login(client, "300300")  # menejer ham o'quv dasturini tuza oladi
    p = Program(name="Kurs V", kind="kurs", price=1)
    db.session.add(p)
    db.session.commit()
    assert "O'quv dasturi" in client.get("/oquv/").get_data(as_text=True)
    client.post(f"/oquv/{p.id}/bolim", data={"title": "Kirish"})
    client.post(f"/oquv/{p.id}/bolim", data={"title": "Amaliyot"})
    m1, m2 = Module.query.order_by(Module.position).all()
    client.post(f"/oquv/bolim/{m2.id}/surish", data={"dir": "up"})
    assert [x.title for x in Module.query.order_by(Module.position).all()] == ["Amaliyot", "Kirish"]
    client.post(f"/oquv/{p.id}/mavzu", data={"module_id": m1.id, "title": "Video 1", "kind": "video", "published": "on"})
    client.post(f"/oquv/{p.id}/mavzu", data={"module_id": m1.id, "title": "Matn", "kind": "matn", "minutes": "7",
                                             "body": "Salom", "published": "on", "free": "on"})
    v, t = Lecture.query.order_by(Lecture.position).all()
    assert t.duration == 420 and t.free and v.published
    # Bunny sozlanmagan — video yuklash rad etiladi; menejer sozlamani o'zgartira olmaydi
    assert client.post(f"/oquv/mavzu/{v.id}/video").status_code == 400
    assert client.post("/oquv/sozla", data={"bunny_library": "42"}).status_code == 403
    rahbar = app.test_client()
    login(rahbar, "100100")
    rahbar.post("/oquv/sozla", data={"bunny_library": "42", "bunny_key": "secret-key", "bunny_token_key": "tk",
                                     "prayer_method": "ISNA", "prayer_asr": "shafii", "app_city": "Samarqand",
                                     "off_bomdod": "3", "off_xufton": "-99"})
    assert F.get_setting("prayer_offsets") == "3,0,0,0,0,-30" and F.get_setting("app_city") == "Samarqand"
    assert "O'quvchi ilovasi" in rahbar.get("/sozlamalar/?bolim=ilova").get_data(as_text=True)
    r = client.post(f"/oquv/mavzu/{v.id}/video").get_json()
    expire = r["headers"]["AuthorizationExpire"]
    assert r["ok"] and r["headers"]["VideoId"] == "vid-123" and r["endpoint"].endswith("/tusupload")
    assert r["headers"]["AuthorizationSignature"] == hashlib_sha256(f"42secret-key{expire}vid-123")
    assert "secret-key" not in str(r)
    s = client.get(f"/oquv/mavzu/{v.id}/holat?uploaded=1").get_json()
    assert s["status"] == "tayyor" and db.session.get(Lecture, v.id).duration == 754
    url = bunny.embed_url("vid-123", now=1000)
    assert url.startswith("https://iframe.mediadelivery.net/embed/42/vid-123?token=")
    assert f"token={hashlib_sha256('tkvid-12315400')}&expires=15400" in url  # 4 soat
    # o'chirish Bunny'dagi videoni ham o'chiradi
    client.post(f"/oquv/mavzu/{v.id}/ochirish")
    assert ("DELETE", "/videos/vid-123", None) in calls and db.session.get(Lecture, v.id) is None
    # o'quvchi kartasidan taklif havolasi
    mem = Member(full_name="Test O'quvchi")
    db.session.add(mem)
    db.session.commit()
    inv = client.post(f"/oquv/taklif/{mem.id}").get_json()
    assert "/app/k/" in inv["link"] and "<svg" in inv["qr"]


def hashlib_sha256(s):
    import hashlib
    return hashlib.sha256(s.encode()).hexdigest()


def test_prayer_api_manifest_and_service_worker(app):
    client = app.test_client()
    d = client.get("/app/api/namoz?city=Buxoro&days=7").get_json()
    assert d["city"] == "Buxoro" and len(d["today"]) == 6 and len(d["week"]) == 7
    assert d["next"]["at"].endswith("+05:00") and 230 < d["qibla"]["deg"] < 245
    d = client.get("/app/api/namoz?lat=21.4&lng=39.8&tz=3").get_json()
    assert d["city"] is None and d["qibla"]["km"] < 10
    assert client.get("/app/api/namoz?lat=abc").get_json()["city"] == "Toshkent"
    man = client.get("/app/manifest.webmanifest")
    assert man.mimetype == "application/manifest+json" and man.get_json()["scope"] == "/app/"
    sw = client.get("/app/sw.js")
    assert sw.mimetype == "application/javascript" and b"caches" in sw.data
    r = client.get("/app/kirish")
    assert "frame-ancestors" in r.headers["Content-Security-Policy"] and "X-Frame-Options" not in r.headers


def test_members_filters_sort_and_export(app):
    from io import BytesIO
    from openpyxl import load_workbook
    from core.attendance import member_code
    from core.models import TelegramLink
    client = app.test_client()
    login(client, "100100")
    kurs = Program(name="Kurs F", kind="kurs", price=1_000_000)
    klub = Program(name="Klub F", kind="klub", price=500_000)
    a = Member(full_name="Ali Valiyev", phone="+998 90 111 22 33", city="Toshkent")
    b = Member(full_name="Bobur Sodiqov", phone="+998 91 444 55 66", city="Samarqand")
    c = Member(full_name="Sardor Karimov", phone="", city="Toshkent")
    db.session.add_all([kurs, klub, a, b, c])
    db.session.commit()
    ea = Enrollment(member_id=a.id, program_id=kurs.id, price=1_000_000, joined_at=date(2026, 1, 1))
    eb = Enrollment(member_id=b.id, program_id=klub.id, price=500_000, joined_at=date(2026, 1, 1),
                    valid_until=date.today() - timedelta(days=2))
    ec = Enrollment(member_id=c.id, program_id=klub.id, price=500_000, joined_at=date(2026, 1, 1),
                    valid_until=date.today() + timedelta(days=3))
    db.session.add_all([ea, eb, ec, TelegramLink(member_id=b.id, chat_id="5")])
    db.session.commit()
    F.record("kirim", 1_000_000, acc("Asosiy kassa"), enrollment_id=ea.id, program_id=kurs.id)
    F.record("kirim", 200_000, acc("Asosiy kassa"), enrollment_id=ec.id, program_id=klub.id)
    db.session.commit()

    def names(**params):
        html = client.get("/azolar", query_string=params).get_data(as_text=True)
        return [n for n in ("Ali Valiyev", "Bobur Sodiqov", "Sardor Karimov") if n in html]

    assert names() == ["Ali Valiyev", "Bobur Sodiqov", "Sardor Karimov"]
    assert names(holat="qarzdor") == ["Bobur Sodiqov", "Sardor Karimov"]
    assert names(holat="tugagan") == ["Bobur Sodiqov"]
    assert names(holat="tugayapti") == ["Sardor Karimov"]
    assert names(dastur=klub.id) == ["Bobur Sodiqov", "Sardor Karimov"]
    assert names(tur="kurs") == ["Ali Valiyev"]
    assert names(shahar="Toshkent", tg="0") == ["Ali Valiyev", "Sardor Karimov"]
    assert names(tg="1") == ["Bobur Sodiqov"]
    assert names(q="914445566") == ["Bobur Sodiqov"]          # telefon — probelsiz raqam bilan
    assert names(q=member_code(c.id)) == ["Sardor Karimov"]   # QR karta kodi bilan
    assert names(q="samarqand") == ["Bobur Sodiqov"]
    html = client.get("/azolar", query_string={"tartib": "qarz"}).get_data(as_text=True)
    assert html.index("Bobur Sodiqov") < html.index("Sardor Karimov") < html.index("Ali Valiyev")
    html = client.get("/azolar", query_string={"dastur": klub.id, "holat": "qarzdor"}).get_data(as_text=True)
    assert "Klub F" in html and "Hammasini tozalash" in html
    # Excel — filtr bilan
    r = client.get("/azolar/excel", query_string={"dastur": klub.id})
    ws = load_workbook(BytesIO(r.data)).active
    assert [ws.cell(i, 1).value for i in (2, 3)] == ["Bobur Sodiqov", "Sardor Karimov"] and ws.cell(4, 1).value == "Jami"
    manager = app.test_client()
    login(manager, "300300")
    assert manager.get("/azolar/excel").status_code == 302


# ── Sun'iy intellekt (Claude javobi soxtalashtirilgan) ──

class _FakeClaude:
    """client.beta.messages.create(...) o'rnini bosadi: navbatdagi javobni qaytaradi, so'rovni yozib qo'yadi."""

    def __init__(self):
        self.replies, self.calls = [], []
        outer = self

        class _Msgs:
            def create(self, **kw):
                outer.calls.append(kw)
                from types import SimpleNamespace as NS
                text = outer.replies.pop(0)
                if not isinstance(text, str):
                    import json as _j
                    text = _j.dumps(text, ensure_ascii=False)
                return NS(content=[NS(type="thinking", thinking=""), NS(type="text", text=text)],
                          stop_reason="end_turn",
                          usage=NS(input_tokens=1200, output_tokens=300, cache_read_input_tokens=0,
                                   cache_creation_input_tokens=0))

            def stream(self, **kw):
                msg = self.create(**kw)

                class _S:
                    def __enter__(self_inner):
                        return self_inner

                    def __exit__(self_inner, *a):
                        return False

                    def get_final_message(self_inner):
                        return msg
                return _S()

        from types import SimpleNamespace as NS
        self.beta = NS(messages=_Msgs())


@pytest.fixture
def claude(app, monkeypatch):
    from core import ai
    fake = _FakeClaude()
    monkeypatch.setattr(ai, "_client", lambda: fake)
    F.set_setting("ai_key", "sk-ant-test")
    db.session.commit()
    return fake


def _staff_link(client, role_code="100100", chat=555):
    from core import staff_bot as S
    from core import telegram as T
    from core.models import User
    user = next(u for u in User.query.all() if u.check_code(role_code))
    tok = S.new_token(user)
    secret = T.webhook_secret()
    hdr = {"X-Telegram-Bot-Api-Secret-Token": secret}
    client.post(f"/telegram/hook/{secret}", json=_upd(chat, f"/start X{tok}"), headers=hdr)
    return user, hdr, f"/telegram/hook/{secret}"


def test_ai_kassa_entry_from_telegram_with_confirmation(app, tg, claude):
    from core.models import AiDraft, AiUsage, StaffLink
    client = app.test_client()
    p = Program(name="Kurs AI", kind="kurs", price=3_000_000)
    m = Member(full_name="Akmal Karimov")
    db.session.add_all([p, m])
    db.session.commit()
    e = Enrollment(member_id=m.id, program_id=p.id, price=3_000_000, joined_at=date(2026, 1, 1))
    db.session.add(e)
    db.session.commit()
    user, hdr, hook = _staff_link(client)
    link = StaffLink.query.filter_by(user_id=user.id).one()
    assert link.active and link.chat_id == "555" and link.token is None
    assert "Ulandi" in _sent(tg)[-1]["text"]
    # matndan yozuv: o'quvchi to'lovi, hisob ko'rsatilmagan
    claude.replies.append({"intent": "yozuv", "reply": "", "entries": [{
        "kind": "kirim", "amount": 1500000, "currency": "UZS", "account_id": 0, "to_account_id": 0,
        "category_id": 0, "enrollment_id": e.id, "counterparty": "", "date": date.today().isoformat(),
        "note": "Kurs to'lovi"}]})
    client.post(hook, json=_upd(555, "Akmaldan 1.5 mln naqd kurs to'lovi oldim"), headers=hdr)
    call = claude.calls[-1]
    assert call["model"] == "claude-opus-5-5" and call["output_config"]["format"]["type"] == "json_schema"
    assert "Akmal Karimov" in call["system"][0]["text"] and "TODAY:" in call["messages"][0]["content"][-1]["text"]
    msg = _sent(tg)[-1]
    assert "1 500 000 so'm" in msg["text"] and "Akmal Karimov" in msg["text"] and "qarzidan" in msg["text"]
    draft = AiDraft.query.one()
    assert msg["reply_markup"]["inline_keyboard"][0][0]["callback_data"] == f"ai:ok:{draft.id}"
    assert Transaction.query.count() == 0  # tasdiqlanmaguncha kassaga tushmaydi
    # «Saqlash» tugmasi
    cq = {"update_id": next(_UPD_IDS), "callback_query": {"id": "q1", "data": f"ai:ok:{draft.id}",
          "message": {"message_id": 7, "chat": {"id": 555, "type": "private"}, "text": msg["text"]}}}
    client.post(hook, json=cq, headers=hdr)
    tx = Transaction.query.one()
    assert (tx.kind, tx.amount, tx.enrollment_id, tx.created_by) == ("kirim", 1_500_000, e.id, user.id)
    assert tx.category_id is not None  # dastur to'lovi kategoriyasi o'zi qo'yildi
    assert any(meth == "editMessageText" for meth, _ in tg)
    cq2 = dict(cq, update_id=next(_UPD_IDS))
    client.post(hook, json=cq2, headers=hdr)  # ikkinchi bosish — takror yozilmaydi
    assert Transaction.query.count() == 1
    assert AiUsage.query.count() == 1 and AiUsage.query.first().cost_usd > 0


def test_ai_kassa_receipt_photo_question_and_limits(app, tg, claude, monkeypatch):
    from core import ai, ai_kassa as K
    from core import telegram as T
    from core.models import AiDraft
    client = app.test_client()
    user, hdr, hook = _staff_link(client, "300300", chat=777)  # menejer
    # menejer chiqim yozolmaydi
    claude.replies.append({"intent": "yozuv", "reply": "", "entries": [{
        "kind": "chiqim", "amount": 250000, "currency": "UZS", "account_id": acc("Asosiy kassa").id,
        "to_account_id": 0, "category_id": 0, "enrollment_id": 0, "counterparty": "Korzinka",
        "date": date.today().isoformat(), "note": "Ofis uchun"}]})
    monkeypatch.setattr(T, "download_file", lambda fid, limit=10_000_000: b"\xff\xd8fakejpeg")
    client.post(hook, json=_upd(777, photo=[{"file_id": "s", "file_size": 100}, {"file_id": "big", "file_size": 2000}],
                                caption="chek"), headers=hdr)
    img = claude.calls[-1]["messages"][0]["content"][0]
    assert img["type"] == "image" and img["source"]["media_type"] == "image/jpeg"
    assert "Menejer faqat kirim" in _sent(tg)[-1]["text"] and AiDraft.query.count() == 0
    # menejerga moliyaviy savol berilmaydi
    claude.replies.append({"intent": "savol", "entries": [], "reply": "Bu oy nega zarar?"})
    client.post(hook, json=_upd(777, "bu oy nega zarar?"), headers=hdr)
    assert "faqat rahbar va buxgalter" in _sent(tg)[-1]["text"]
    # rahbar savoli — raqamlar dasturda hisoblanadi, javob xavfsiz HTML
    boss, hdr2, _ = _staff_link(client, "100100", chat=888)
    claude.replies += [{"intent": "savol", "entries": [], "reply": "Bu oy nega zarar?"},
                       "Chunki <b>xarajat</b> oshdi <script>x</script>"]
    client.post(hook, json=_upd(888, "bu oy nega zarar?"), headers=hdr2)
    assert "this_month_so_far" in claude.calls[-1]["messages"][0]["content"]
    assert _sent(tg)[-1]["text"] == "Chunki <b>xarajat</b> oshdi &lt;script&gt;x&lt;/script&gt;"
    # tozalash: dollar → so'm, noto'g'ri kategoriya, kelajakdagi sana
    cat_out = Category.query.filter_by(kind="chiqim").first()
    e = K.clean({"kind": "kirim", "amount": 100, "currency": "USD", "account_id": acc("Asosiy kassa").id,
                 "to_account_id": 0, "category_id": cat_out.id, "enrollment_id": 0, "counterparty": "X",
                 "date": "2030-01-01", "note": ""})
    assert e["amount"] == round(100 * F.usd_rate(), 2) and e["category_id"] is None and e["date"] == date.today().isoformat()
    assert K.clean({"kind": "kirim", "amount": 0, "currency": "UZS", "account_id": 0}) is None
    # kunlik limit
    F.set_setting("ai_daily_limit", "1")
    db.session.commit()
    with pytest.raises(ai.AiError):
        ai.ask("s", "x", feature="test")


def test_ai_weekly_report_dashboard_and_settings(app, tg, claude, monkeypatch):
    from datetime import datetime
    from core import ai, ai_report as R
    from core.models import AiReport
    client = app.test_client()
    login(client, "100100")
    user, hdr, hook = _staff_link(client)
    p = Program(name="Klub R", kind="klub", price=1_000_000)
    m = Member(full_name="Qarzdor Ali")
    db.session.add_all([p, m])
    db.session.commit()
    db.session.add(Enrollment(member_id=m.id, program_id=p.id, price=1_000_000, joined_at=date(2026, 1, 1)))
    F.record("chiqim", 400_000, acc("Asosiy kassa"), category_id=Category.query.filter_by(kind="chiqim").first().id)
    db.session.commit()
    snap = R.snapshot()
    assert snap["debtors_count"] == 1 and snap["top_debtors"][0]["name"] == "Qarzdor Ali"
    assert snap["this_month_so_far"]["expense"] == 400_000
    # yakshanba 10:00 (standart) — hisobot bir marta tuziladi va rahbarga yuboriladi
    sunday = datetime(2026, 10, 18, 10, 5)
    claude.replies.append("📊 <b>Haftalik hisobot</b>\nQarzdor Ali — 1 mln")
    rep = R.run_due(sunday)
    assert rep and rep.period == "2026-W42" and "Qarzdor Ali" in claude.calls[-1]["messages"][0]["content"]
    assert _sent(tg)[-1]["chat_id"] == "555" and "Haftalik hisobot" in _sent(tg)[-1]["text"]
    assert R.run_due(sunday) is None and AiReport.query.count() == 1
    assert R.run_due(datetime(2026, 10, 19, 10, 0)) is None  # dushanba — hisobot kuni emas
    # bosh sahifada karta faqat hisobot kuni chiqadi
    F.set_setting("ai_weekly_day", str((date.today().weekday() + 1) % 7))
    db.session.commit()
    assert "Haftalik tahlil" not in client.get("/").get_data(as_text=True)
    F.set_setting("ai_weekly_day", str(date.today().weekday()))
    db.session.commit()
    page = client.get("/").get_data(as_text=True)
    assert "Haftalik tahlil" in page and "Haftalik hisobot" in page
    # qo'lda yangilash
    claude.replies.append("Yangi tahlil")
    client.post("/ai/hisobot", data={"next": "/"})
    assert R.latest().text == "Yangi tahlil"
    # sozlamalar: kalit tekshiriladi
    monkeypatch.setattr(ai, "check_key", lambda k: (k == "sk-ant-good", "Kalit noto'g'ri."))
    client.post("/ai/sozla", data={"ai_key": "bad", "ai_weekly": "on", "ai_weekly_day": "4", "ai_weekly_hour": "8",
                                   "ai_daily_limit": "50"})
    assert F.get_setting("ai_key") == "sk-ant-test"
    client.post("/ai/sozla", data={"ai_key": "sk-ant-good", "ai_weekly_day": "4", "ai_weekly_hour": "8",
                                   "ai_daily_limit": "50"})
    assert F.get_setting("ai_key") == "sk-ant-good" and F.get_setting("ai_weekly") == "0"
    assert F.get_setting("ai_weekly_day") == "4"
    F.set_setting("ai_weekly_day", "6")
    db.session.commit()
    page = client.get("/sozlamalar/?bolim=ai").get_data(as_text=True)
    assert "Anthropic API kaliti" in page and "ulangan" in page
    assert "Mening Telegram" in client.get("/ai/telegram").get_data(as_text=True)
    manager = app.test_client()
    login(manager, "300300")
    assert manager.post("/ai/hisobot").status_code == 403
    assert manager.get("/ai/telegram").status_code == 200


# ── Shartnoma tekshiruvi ──

def _docx(text):
    import io
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        paras = "".join(f"<w:p><w:r><w:t>{t}</w:t></w:r></w:p>" for t in text.split("\n"))
        z.writestr("word/document.xml", f'<w:document xmlns:w="x"><w:body>{paras}</w:body></w:document>')
    return buf.getvalue()


def _kb(title="Kechikish jarimasi joiz emas", topic="umumiy", body="Kreditor foydasiga jarima — ribo."):
    from core.models import KnowledgeDoc
    d = KnowledgeDoc(title=title, topic=topic, kind="qoida", body=body, source="Ustoz, 3-dars")
    db.session.add(d)
    db.session.commit()
    return d


def test_knowledge_base_crud_and_text_import(app):
    from io import BytesIO
    from core import contract as C
    from core.models import KnowledgeDoc
    client = app.test_client()
    login(client, "300300")
    r = client.post("/shartnoma/bilimlar/saqlash", data={
        "title": "Murobahada egalik", "kind": "ustoz", "topic": "murobaha", "source": "Ustoz Abdulloh",
        "body": "Sotuvchi avval egalik qiladi.", "file": (BytesIO(_docx("Birinchi qator\nIkkinchi qator")), "q.docx")},
        content_type="multipart/form-data")
    assert r.status_code == 302
    d = KnowledgeDoc.query.one()
    assert d.body.startswith("Sotuvchi avval") and "Ikkinchi qator" in d.body and d.topic == "murobaha"
    page = client.get("/shartnoma/bilimlar?mavzu=murobaha").get_data(as_text=True)
    assert "Murobahada egalik" in page and "Ustoz Abdulloh" in page
    assert "Murobahada egalik" not in client.get("/shartnoma/bilimlar?mavzu=ijara").get_data(as_text=True)
    assert C.text_from_upload("a.txt", "Salom".encode()) == "Salom"
    # katta baza: shartnoma mavzusi + umumiy materiallar qoladi
    _kb("Umumiy qoida", "umumiy", "x" * 50)
    _kb("Ijara qoidasi", "ijara", "y" * 50)
    import core.contract as CM
    old = CM.KB_CHAR_BUDGET
    CM.KB_CHAR_BUDGET = 400
    try:
        text, ids, trimmed = C.knowledge(C.detect_topics("Murobaha shartnomasi, ustama 20%"))
    finally:
        CM.KB_CHAR_BUDGET = old
    titles = {db.session.get(KnowledgeDoc, i).title for i in ids}
    assert trimmed and titles == {"Murobahada egalik", "Umumiy qoida"}
    client.post(f"/shartnoma/bilimlar/{d.id}/ochirish")
    assert KnowledgeDoc.query.count() == 2


def _result(src_ids):
    return {"contract_type": "Oldi-sotdi (nasiya)", "summary": "Tovar bo'lib to'lashga sotiladi.",
            "verdict": "mos", "verdict_reason": "Asosan to'g'ri, bir band tuzatilishi kerak.",
            "findings": [
                {"severity": "orta", "category": "jarima", "clause": "kechikkan har kun uchun 0.5% penya",
                 "issue": "Kechikish jarimasi sotuvchi foydasiga", "explanation": "Bazaga ko'ra ribo.",
                 "sources": src_ids, "suggestion": "Kechikish uchun summa xayriyaga yo'naltiriladi."},
                {"severity": "jiddiy", "category": "boshqa", "clause": "", "issue": "Sug'urta sharti",
                 "explanation": "?", "sources": [9999], "suggestion": ""}],
            "good_points": ["Narx aniq"], "missing": ["Tovar egaligi o'tish sanasi"], "questions_for_ustoz": []}


def test_contract_check_staff_flow_and_review(app, tg, claude):
    from core.models import ContractCheck, Notification, TelegramLink
    client = app.test_client()
    login(client, "100100")
    # baza bo'sh — tekshiruv xato bilan tugaydi
    r = client.post("/shartnoma/yangi", data={"title": "Bo'sh", "text": "Shartnoma matni"})
    c0 = ContractCheck.query.one()
    assert c0.status == "xato" and "bo'sh" in c0.error
    kb = _kb()
    claude.replies.append(_result([kb.id]))
    r = client.post("/shartnoma/yangi", data={"title": "Yetkazib berish",
                                              "text": "Xaridor kechikkan har kun uchun 0.5% penya to'laydi."})
    c = ContractCheck.query.order_by(ContractCheck.id.desc()).first()
    assert r.headers["Location"].endswith(f"/shartnoma/{c.id}")
    call = claude.calls[-1]
    sys_text = call["system"][0]["text"]
    assert "ONLY against the centre's KNOWLEDGE BASE" in sys_text and f'<item id="{kb.id}"' in sys_text
    assert call["system"][0]["cache_control"]["ttl"] == "1h" and call["output_config"]["effort"] == "high"
    assert "0.5% penya" in call["messages"][0]["content"][-1]["text"]
    assert c.status == "tayyor" and c.verdict == "shartli"  # «mos», lekin tuzatish bor — shartli
    import json as _j
    res = _j.loads(c.result)
    assert len(res["findings"]) == 1 and "bazada asos topilmadi" in res["questions_for_ustoz"][0]
    page = client.get(f"/shartnoma/{c.id}").get_data(as_text=True)
    assert "Kechikish jarimasi sotuvchi foydasiga" in page and kb.title in page and "Ustoz xulosasi" in page
    assert client.get(f"/shartnoma/{c.id}/holat").get_json()["status"] == "tayyor"
    # o'quvchidan kelgan tekshiruvga ustoz xulosasi — Telegram'ga xabar
    m = Member(full_name="Ali Valiyev")
    db.session.add(m)
    db.session.commit()
    db.session.add(TelegramLink(member_id=m.id, chat_id="321"))
    c.member_id = m.id
    db.session.commit()
    client.post(f"/shartnoma/{c.id}/xulosa", data={"verdict": "mos_emas", "text": "7-bandni o'zgartiring", "notify": "on"})
    c = db.session.get(ContractCheck, c.id)
    assert c.review_verdict == "mos_emas" and c.reviewer.role == "rahbar"
    assert "Mos emas" in _sent(tg)[-1]["text"] and _sent(tg)[-1]["chat_id"] == "321"
    assert Notification.query.filter_by(member_id=m.id).count() == 1
    assert "Hammasi" in client.get("/shartnoma/?holat=kutilmoqda").get_data(as_text=True)


def test_contract_check_student_app(app, claude):
    from io import BytesIO
    from core import student as S
    from core.models import ContractCheck
    kb = _kb()
    m1, m2 = Member(full_name="Bir"), Member(full_name="Ikki")
    db.session.add_all([m1, m2])
    db.session.commit()
    c1, c2 = app.test_client(), app.test_client()
    c1.get(f"/app/k/{S.issue_invite(m1)}")
    c2.get(f"/app/k/{S.issue_invite(m2)}")
    assert "Tekshirish" in c1.get("/app/shartnoma").get_data(as_text=True)
    F.set_setting("contract_student_daily", "1")
    db.session.commit()
    claude.replies.append(_result([kb.id]))
    r = c1.post("/app/shartnoma", data={"title": "Ijara", "files": (BytesIO(b"%PDF-1.4 fake"), "ijara.pdf")},
                content_type="multipart/form-data")
    check = ContractCheck.query.one()
    assert check.member_id == m1.id and check.source_kind == "pdf" and check.status == "tayyor"
    block = claude.calls[-1]["messages"][0]["content"][0]
    assert block["type"] == "document" and block["source"]["media_type"] == "application/pdf"
    page = c1.get(f"/app/shartnoma/{check.id}").get_data(as_text=True)
    assert "Kechikish jarimasi sotuvchi foydasiga" in page and "fatvo emas" in page
    assert c2.get(f"/app/shartnoma/{check.id}").status_code == 404  # boshqa o'quvchiniki
    # kunlik limit
    r = c1.post("/app/shartnoma", data={"text": "matn"}, follow_redirects=True)
    assert ContractCheck.query.count() == 1
    # yopiq bo'lsa
    F.set_setting("contract_students", "0")
    db.session.commit()
    assert "hozircha yopiq" in c2.get("/app/shartnoma").get_data(as_text=True)
    assert c2.post("/app/shartnoma", data={"text": "matn"}).status_code == 302 and ContractCheck.query.count() == 1
